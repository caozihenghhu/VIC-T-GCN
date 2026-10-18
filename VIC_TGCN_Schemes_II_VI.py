"""VIC-T-GCN comparison experiments: Schemes II-VI."""
from __future__ import annotations

import argparse
import copy
import hashlib
import io
import json
import os
import random
import traceback
from datetime import datetime, timezone
from pathlib import Path

import numpy as np
import pandas as pd

PACKAGE_ROOT = Path(__file__).resolve().parent

COMPARISON_OUTPUT = str(PACKAGE_ROOT / "results/schemes_II_VI")
SCHEMES = [2, 3, 4, 5]
SEEDS = [42]
SELECTED_BASINS = None
BASINS = [{
    "name": "danba",
    "runoff_excel": str(PACKAGE_ROOT / "data/danba/runoff.xlsx"),
    "baseflow_excel": str(PACKAGE_ROOT / "data/danba/baseflow.xlsx"),
    "node_table": str(PACKAGE_ROOT / "data/danba/nodes.xlsx"),
    "flow_dir_surface": str(PACKAGE_ROOT / "data/danba/surface_direction.txt"),
    "flow_dir_baseflow": str(PACKAGE_ROOT / "data/danba/baseflow_direction.txt"),
    "output_dir": str(PACKAGE_ROOT / "results/scheme_I"),
}]

DEFAULT_TRAINING = {
    "window": 12, "target_lead_days": 1,
    "train_fraction": 0.65, "val_fraction": 0.10,
    "train_end": None, "val_end": None, "boundary_context": "past",
    "gcn_hidden": 32, "lstm_hidden": 128, "dropout": 0.3,
    "learning_rate": 0.001, "weight_decay": 1e-4, "batch_size": 8,
    "epochs": 200, "early_stop_patience": 15, "lr_patience": 5,
    "seed": 42, "deterministic": True, "device": "auto", "make_plots": True,
}
TRAINING = dict(DEFAULT_TRAINING)

DEFAULT_GRAPHS = {
    "surface": {"encoding": "north_clockwise_1to8", "crs": "EPSG:4326", "edge_rule": "adjacent"},
    "baseflow": {"encoding": "north_clockwise_1to8", "crs": "EPSG:4326", "edge_rule": "adjacent"},
}

DIRECTION_MAPS = {
    "north_clockwise_1to8": {1:(-1,0),2:(-1,1),3:(0,1),4:(1,1),5:(1,0),6:(1,-1),7:(0,-1),8:(-1,-1)},
    "east_clockwise_1to8": {1:(0,1),2:(1,1),3:(1,0),4:(1,-1),5:(0,-1),6:(-1,-1),7:(-1,0),8:(-1,1)},
    "esri_d8": {1:(0,1),2:(1,1),4:(1,0),8:(1,-1),16:(0,-1),32:(-1,-1),64:(-1,0),128:(-1,1)},
}

def merged_graph_config(overrides=None):

    config = copy.deepcopy(DEFAULT_GRAPHS)
    if overrides is not None:
        if not isinstance(overrides,dict) or set(overrides)-set(config):
            raise ValueError("graphs may contain surface and baseflow settings only.")
        for graph, settings in overrides.items():
            if not isinstance(settings,dict) or set(settings)-set(config[graph]):
                raise ValueError(f"Unknown graph setting in {graph}; use encoding/crs/edge_rule.")
            config[graph].update(settings)
    return config

def file_sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for block in iter(lambda: f.read(1024*1024), b""):
            h.update(block)
    return h.hexdigest()

def read_table(path):
    path = Path(path)
    df = pd.read_excel(path) if path.suffix.lower() in (".xlsx", ".xls") else pd.read_csv(path, encoding="utf-8-sig")
    df.columns = [str(c).strip() for c in df.columns]
    if df.columns.duplicated().any():
        raise ValueError(f"Duplicate column names: {path}")
    return df

def parse_dates(values):

    v = pd.Series(values)
    if v.isna().any():
        raise ValueError("The date column contains missing values.")
    text = v.astype(str).str.strip()
    if text.str.fullmatch(r"\d{8}").all():
        dates = pd.to_datetime(text, format="%Y%m%d", errors="raise")
    elif pd.api.types.is_numeric_dtype(v):
        if not ((v >= 20000) & (v <= 80000) & (v == np.floor(v))).all():
            raise ValueError("The numeric first column is not a recognized Excel date; provide valid dates.")
        dates = pd.to_datetime(v, unit="D", origin="1899-12-30")
    else:
        dates = pd.to_datetime(v, errors="raise")
    dates = pd.DatetimeIndex(dates)
    if dates.tz is not None or not dates.equals(dates.normalize()):
        raise ValueError("Dates must be timezone-naive daily dates without time components.")
    if dates.has_duplicates:
        raise ValueError("Duplicate dates.")
    return dates

def read_daily(path):
    df = read_table(path)
    if df.shape[1] < 3 or len(df) == 0:
        raise ValueError(f"{path} requires dates, outlet observations, and at least one node input column.")
    df.index = parse_dates(df.iloc[:, 0])
    df = df.sort_index()
    if not df.index.equals(pd.date_range(df.index[0], df.index[-1], freq="D")):
        raise ValueError(f"{path} has nonconsecutive dates; missing days must not be treated as consecutive days.")
    return df

def validate_node_mapping(frame):

    required = ["node_index","node_id","lat","lon","surface_column","baseflow_column"]
    if not set(required).issubset(frame.columns):
        raise ValueError(f"The node table must contain columns: {required}")
    if len(frame) == 0 or frame[required].isna().any().any():
        raise ValueError("The node table is empty or required fields are missing.")
    nodes = frame[required].copy()
    for key in ("node_id","surface_column","baseflow_column"):
        nodes[key] = nodes[key].map(lambda v: str(int(v)) if isinstance(v,(float,np.floating)) and np.isfinite(v) and v.is_integer() else str(v).strip())
        if nodes[key].eq("").any() or nodes[key].duplicated().any():
            raise ValueError(f"Node table field {key} contains missing or duplicate values.")
    index = pd.to_numeric(nodes.node_index,errors="raise").to_numpy(float)
    if not np.isfinite(index).all() or not np.array_equal(np.sort(index),np.arange(len(nodes))):
        raise ValueError("node_index must contain unique consecutive integers from 0 to N-1.")
    nodes["node_index"] = index.astype(int)
    for key in ("lat","lon"):
        nodes[key] = pd.to_numeric(nodes[key],errors="raise").astype(float)
    if not np.isfinite(nodes[["lat","lon"]].to_numpy()).all():
        raise ValueError("Node coordinates contain missing or infinite values.")
    if not (nodes.lat.between(-90,90).all() and nodes.lon.between(-180,180).all()):
        raise ValueError("Node coordinates must be EPSG:4326 longitude/latitude, not projected coordinates.")
    if nodes[["lat","lon"]].duplicated().any():
        raise ValueError("Multiple nodes share the same coordinates.")
    nodes["coordinate_source"] = "node_table"
    nodes["mapping_rule"] = "explicit_node_index_and_named_input_columns"
    return nodes.sort_values("node_index").reset_index(drop=True)

def load_basin_inputs(basin):
    s, b = read_daily(basin["runoff_excel"]), read_daily(basin["baseflow_excel"])
    if not s.index.equals(b.index):
        raise ValueError("Runoff and baseflow dates differ; row-based concatenation is not allowed.")
    node_path = Path(basin["node_table"])
    if node_path.suffix.lower() in (".xlsx",".xls"):
        frame = pd.read_excel(node_path,sheet_name="Nodes",
                              dtype={"node_id":str,"surface_column":str,"baseflow_column":str})
    else:
        frame = pd.read_csv(node_path,encoding="utf-8-sig",
                            dtype={"node_id":str,"surface_column":str,"baseflow_column":str})
    nodes = validate_node_mapping(frame)
    for label, table, key in (("runoff",s,"surface_column"),("baseflow",b,"baseflow_column")):
        actual = list(table.columns[2:])
        expected = nodes[key].tolist()
        if len(actual) != len(expected) or set(actual) != set(expected):
            raise ValueError(f"{label} node columns do not match the node table; positional matching is not allowed.")
    xs = s.loc[:,nodes.surface_column].to_numpy(dtype=np.float64)
    xb = b.loc[:,nodes.baseflow_column].to_numpy(dtype=np.float64)
    y = s.iloc[:,1].to_numpy(dtype=np.float64)
    base_y = b.iloc[:,1].to_numpy(dtype=np.float64)
    if not np.array_equal(y,base_y):
        raise ValueError("Observed discharge differs between the runoff and baseflow tables.")
    for label, a in [("runoff",xs),("baseflow",xb),("observed discharge",y)]:
        if not np.isfinite(a).all() or np.any(a < 0):
            raise ValueError(f"{label} contains missing, infinite, or negative values; full-period interpolation is not performed.")
    if np.array_equal(xs,xb):
        raise ValueError("Runoff and baseflow are identical; check whether the same file was loaded twice.")
    print(f"Node table loaded: {len(nodes)} nodes; ordered by node_index and matched to both inputs by column name.",flush=True)
    return s.index, np.stack([xs,xb],axis=-1), y, nodes

def make_windows(dates, cfg):

    n, window, lead = len(dates), int(cfg["window"]), int(cfg["target_lead_days"])
    if window < 1 or lead < 0:
        raise ValueError("window must be positive and target_lead_days must be nonnegative.")
    if bool(cfg.get("train_end")) != bool(cfg.get("val_end")):
        raise ValueError("Specify both train_end and val_end, or set both to null.")
    if cfg.get("train_end"):
        t1, t2 = pd.Timestamp(cfg["train_end"]), pd.Timestamp(cfg["val_end"])
        if t1 not in dates or t2 not in dates or t1 >= t2:
            raise ValueError("Split dates must exist and train_end must precede val_end.")
        b1, b2 = dates.get_loc(t1)+1, dates.get_loc(t2)+1
    else:
        f1, f2 = float(cfg["train_fraction"]), float(cfg["val_fraction"])
        if not (0 < f1 < 1 and 0 < f2 < 1 and f1+f2 < 1):
            raise ValueError("Invalid training/validation fractions.")
        b1, b2 = int(n*f1), int(n*(f1+f2))
    if not 0 < b1 < b2 < n:
        raise ValueError("Empty or invalid data split.")
    target = np.arange(window-1+lead, n)
    last, first = target-lead, target-lead-window+1
    split = np.where(target < b1, "train", np.where(target < b2, "val", "test"))
    keep = np.ones(len(target), dtype=bool)
    if cfg["boundary_context"] == "within_split":
        keep = first >= np.where(target < b1, 0, np.where(target < b2, b1, b2))
    elif cfg["boundary_context"] != "past":
        raise ValueError("boundary_context must be past or within_split.")
    table = pd.DataFrame({
        "target_index": target[keep], "input_start_index": first[keep], "input_end_index": last[keep],
        "split": split[keep], "target_date": dates[target[keep]].strftime("%Y-%m-%d"),
        "input_start_date": dates[first[keep]].strftime("%Y-%m-%d"),
        "input_end_date": dates[last[keep]].strftime("%Y-%m-%d")})
    for label in ("train", "val", "test"):
        if not (table.split == label).any():
            raise ValueError(f"{label} has no valid windows.")
    return table, {"train_end": str(dates[b1-1].date()), "val_end": str(dates[b2-1].date()),
                   "test_start": str(dates[b2].date()), "split_basis": "raw_daily_dates"}

def numpy_window(x, row):

    return np.ascontiguousarray(x[int(row.input_start_index):int(row.input_end_index)+1].transpose(1,0,2))

def read_ascii_grid(path):
    try:
        text = Path(path).read_text(encoding="utf-8-sig")
    except UnicodeDecodeError:
        text = Path(path).read_text(encoding="gbk")
    lines, header = text.splitlines(), {}
    known = {"ncols","nrows","xllcorner","yllcorner","xllcenter","yllcenter","cellsize","nodata_value"}
    for start,line in enumerate(lines):
        parts = line.split()
        if not parts:
            continue
        key = parts[0].lower()
        if key not in known:
            break
        if len(parts) != 2 or key in header:
            raise ValueError(f"Invalid ASCII header: {line}")
        header[key] = float(parts[1])
    else:
        raise ValueError("The direction file contains no raster data.")
    for k in ("nrows","ncols","cellsize","nodata_value"):
        if k not in header or not np.isfinite(header[k]):
            raise ValueError(f"The ASCII header is missing a valid field: {k}")
    nr,nc,cs = int(header["nrows"]),int(header["ncols"]),header["cellsize"]
    if nr != header["nrows"] or nc != header["ncols"] or min(nr,nc,cs) <= 0:
        raise ValueError("Invalid raster dimensions or cellsize.")
    origin = []
    for ax in ("x","y"):
        corner,center = ax+"llcorner",ax+"llcenter"
        if (corner in header) == (center in header):
            raise ValueError(f"Specify exactly one of {corner} or {center}.")
        origin.append(header[corner] if corner in header else header[center]-cs/2)
    grid = np.loadtxt(io.StringIO("\n".join(lines[start:])), ndmin=2)
    if grid.shape != (nr,nc) or not np.isfinite(grid).all() or not np.isfinite(origin).all():
        raise ValueError("Invalid raster shape or values.")
    if not np.all(grid == np.rint(grid)):
        raise ValueError("Raster directions are not integers and cannot be used as D8 codes.")
    return grid, origin[0], origin[1], cs, header["nodata_value"]

def resolve_direction_mapping(encoding):
    if isinstance(encoding,str) and encoding in DIRECTION_MAPS:
        return DIRECTION_MAPS[encoding]
    if isinstance(encoding,dict):
        mapping = {int(k):tuple(int(v) for v in d) for k,d in encoding.items()}
        if len(mapping) != 8 or set(mapping.values()) != set(DIRECTION_MAPS["esri_d8"].values()) or 0 in mapping:
            raise ValueError("Custom codes must map uniquely to the eight neighbors; 0 is reserved for no direction.")
        return mapping
    raise ValueError("Unknown direction encoding; use a named DIRECTION_MAPS entry or an explicit mapping.")

def graph_has_cycle(edges, n):

    nxt = dict(edges.T.tolist())
    done = set()
    for start in range(n):
        path, node = set(), start
        while node in nxt and node not in done:
            if node in path:
                return True
            path.add(node)
            node = nxt[node]
        done.update(path)
    return False

def build_flow_graph(path, nodes, encoding="north_clockwise_1to8", crs="EPSG:4326", edge_rule="adjacent"):
    if crs != "EPSG:4326":
        raise ValueError("The ASCII raster and lat/lon must use EPSG:4326; reproject projected rasters first.")
    mapping = resolve_direction_mapping(encoding)
    if edge_rule not in ("adjacent","trace"):
        raise ValueError("edge_rule must be adjacent or trace.")
    grid,xll,yll,cs,nodata = read_ascii_grid(path)
    nr,nc = grid.shape
    bad = sorted(set(np.unique(grid))-set(mapping)-{0,nodata})
    if bad:
        raise ValueError(f"Raster contains undefined direction codes: {bad[:12]}; check the D8 convention and depression codes.")
    r = np.floor((yll+nr*cs-nodes.lat.to_numpy(float))/cs).astype(int)
    c = np.floor((nodes.lon.to_numpy(float)-xll)/cs).astype(int)
    if np.any((r<0)|(r>=nr)|(c<0)|(c>=nc)):
        raise ValueError("Nodes fall outside the raster; check CRS, coordinate order, and spatial coverage.")
    rc = list(zip(r.tolist(),c.tolist()))
    if len(set(rc)) != len(rc):
        raise ValueError("Multiple nodes map to the same raster cell; check resolution rather than overwriting node IDs.")

    missing_direction = grid[r,c] == nodata
    lookup, records = {v:i for i,v in enumerate(rc)}, []
    for i,(r0,c0) in enumerate(rc):
        rr,cc,seen,steps,dst = r0,c0,{(r0,c0)},0,None
        while True:
            code = int(grid[rr,cc])
            if code == nodata:
                status = "source_nodata" if steps == 0 else "path_reaches_nodata"
                break
            if code not in mapping:
                status = "no_direction"
                break
            dr,dc = mapping[code]
            rr,cc,steps = rr+dr,cc+dc,steps+1
            if not (0<=rr<nr and 0<=cc<nc):
                status = "leaves_raster"
                break
            if (rr,cc) in seen:
                status = "flow_cycle"
                break
            if (rr,cc) in lookup:
                dst,status = lookup[(rr,cc)],"edge"
                break
            if edge_rule == "adjacent":
                status = "downstream_cell_has_no_node"
                break
            seen.add((rr,cc))
        records.append({"source_index":i,"source_node":str(nodes.iloc[i].node_id),
                        "source_row":r0,"source_col":c0,"source_code":int(grid[r0,c0]),
                        "source_lat":float(nodes.iloc[i].lat),"source_lon":float(nodes.iloc[i].lon),
                        "source_direction_missing":bool(missing_direction[i]),
                        "target_index":dst,"target_node":None if dst is None else str(nodes.iloc[dst].node_id),
                        "end_row":rr,"end_col":cc,"raster_steps":steps,"status":status})
    audit = pd.DataFrame(records)
    pairs = [(v["source_index"],v["target_index"]) for v in records if v["status"]=="edge"]
    edges = np.array(pairs,dtype=np.int64).T if pairs else np.empty((2,0),dtype=np.int64)
    stats = {"nodes":len(nodes),"edges":int(edges.shape[1]),"isolated_nodes":len(nodes)-len(set(edges.ravel())),
             "has_cycle":graph_has_cycle(edges,len(nodes)),"edge_rule":edge_rule,"cellsize":cs,
             "nodata_nodes":int(missing_direction.sum()),
             "nodata_fraction":float(missing_direction.mean()),
             "nodata_policy":"retain_nodes_no_unknown_outgoing_edges_keep_supported_incoming_edges",
             "routing_direction_coverage_complete":not bool(missing_direction.any()),
             "status_counts":{str(k):int(v) for k,v in audit.status.value_counts().items()},
             "encoding":{str(k):list(v) for k,v in mapping.items()},"coordinate_system_assumption":crs}
    return edges,audit,stats

def torch_dependencies():
    try:
        import torch
        from torch_geometric.nn import GCNConv
        import torch_geometric
    except ImportError as exc:
        raise RuntimeError("Training requires compatible torch and torch_geometric installations; install torch, torch-geometric, numpy, pandas, scipy, openpyxl, and matplotlib.") from exc
    return torch,GCNConv,torch_geometric

def set_seed(seed, deterministic, torch):
    if deterministic:
        os.environ.setdefault("CUBLAS_WORKSPACE_CONFIG",":4096:8")
    random.seed(seed)
    np.random.seed(seed)
    torch.manual_seed(seed)
    if torch.cuda.is_available():
        torch.cuda.manual_seed_all(seed)
    torch.backends.cudnn.benchmark = False
    torch.backends.cudnn.deterministic = bool(deterministic)

    torch.use_deterministic_algorithms(bool(deterministic))

def calc_metrics(obs,pred):
    obs,pred = np.asarray(obs,float).reshape(-1),np.asarray(pred,float).reshape(-1)
    if obs.shape!=pred.shape or not len(obs) or not np.isfinite([obs,pred]).all():
        raise ValueError("Metrics require finite observations and predictions of equal length, paired by day.")
    mse,var,mean = float(np.mean((pred-obs)**2)),float(np.var(obs)),float(np.mean(obs))
    nse = 1-mse/var if var>0 else np.nan
    kge = np.nan
    if len(obs)>1 and var>0 and mean!=0 and np.std(pred)>0:
        r = float(np.corrcoef(obs,pred)[0,1])
        kge = 1-np.sqrt((r-1)**2+(np.std(pred)/np.std(obs)-1)**2+(np.mean(pred)/mean-1)**2)
    return {"NSE":nse,"KGE":float(kge),"MAE":float(np.mean(np.abs(pred-obs))),
            "RMSE":float(np.sqrt(mse))}

def validate_training_config(cfg):
    for k in ("window","batch_size","epochs","early_stop_patience","gcn_hidden","lstm_hidden"):
        if not isinstance(cfg[k],int) or cfg[k]<1:
            raise ValueError(f"{k} must be a positive integer.")
    for k in ("target_lead_days","lr_patience","seed"):
        if not isinstance(cfg[k],int) or cfg[k]<0:
            raise ValueError(f"{k} must be a nonnegative integer.")
    if not (0<=cfg["dropout"]<1 and cfg["learning_rate"]>0 and cfg["weight_decay"]>=0):
        raise ValueError("Invalid dropout, learning_rate, or weight_decay.")

def read_prediction_table(path, feature_columns):

    df = read_table(path)
    if len(df)==0 or len(df.columns)<2:
        raise ValueError(f"Prediction input is empty or lacks node columns: {path}")
    date_column = df.columns[0]
    if date_column in feature_columns:
        raise ValueError("The first prediction-input column must contain dates, not node features.")
    missing = set(feature_columns)-set(df.columns)
    if missing:
        raise ValueError(f"{path} is missing training node columns: {sorted(missing)[:10]}; positional matching is not allowed.")
    df.index = parse_dates(df.iloc[:,0])
    df = df.sort_index()
    if not df.index.equals(pd.date_range(df.index[0],df.index[-1],freq="D")):
        raise ValueError(f"Nonconsecutive prediction input dates: {path}")
    features = df.loc[:,feature_columns].to_numpy(dtype=np.float64)
    if not np.isfinite(features).all() or np.any(features<0):
        raise ValueError(f"Prediction features contain missing, infinite, or negative values: {path}")
    return df,features

VERSION = '1.1.1_input_sync'

GRAPH_CONFIG = copy.deepcopy(DEFAULT_GRAPHS)

STATION_ORDER = ['ribu', 'yilinchang', 'dajin', 'danba', 'xiaojin']
STATION_EDGES = [('ribu','dajin'), ('yilinchang','dajin'), ('dajin','danba'), ('xiaojin','danba')]
STATION_PCA_DIM = 15

def write_json(path, data):
    def clean(v):
        if isinstance(v,dict): return {str(k):clean(x) for k,x in v.items()}
        if isinstance(v,(list,tuple)): return [clean(x) for x in v]
        if isinstance(v,np.generic): return clean(v.item())
        if isinstance(v,float) and not np.isfinite(v): return None
        if isinstance(v,Path): return str(v)
        return v
    path=Path(path); temp=path.with_suffix(path.suffix+'.tmp')
    temp.write_text(json.dumps(clean(data),ensure_ascii=False,indent=2,allow_nan=False),encoding='utf-8')
    os.replace(temp,path)

def scheme_definition(scheme):
    return {
        1: {'channels':[0,1],'graphs':['surface','baseflow'],'branch_inputs':[1,1],'outputs':'one_outlet'},
        2: {'channels':[0,1],'graphs':['surface','surface'],'branch_inputs':[1,1],'outputs':'one_outlet'},
        3: {'channels':[0],'graphs':['surface'],'branch_inputs':[1],'outputs':'one_outlet'},
        4: {'channels':[1],'graphs':['baseflow'],'branch_inputs':[1],'outputs':'one_outlet'},
        5: {'channels':[1],'graphs':['surface'],'branch_inputs':[1],'outputs':'one_outlet'},
        6: {'channels':[0,1],'graphs':['station'],'branch_inputs':[30],'outputs':'all_stations'},
    }[int(scheme)]

def build_scheme_model(n_nodes, cfg, scheme, n_outputs=1):
    torch,GCNConv,_=torch_dependencies()
    spec=scheme_definition(scheme)
    class SchemeModel(torch.nn.Module):
        def __init__(self):
            super().__init__()
            self.n_nodes=n_nodes; self.scheme=int(scheme); self.n_outputs=n_outputs
            h=int(cfg['gcn_hidden'])
            self.first=torch.nn.ModuleList([GCNConv(c,h,flow='source_to_target') for c in spec['branch_inputs']])
            self.second=torch.nn.ModuleList([GCNConv(h,h,flow='source_to_target') for _ in spec['branch_inputs']])
            self.dropout=torch.nn.Dropout(float(cfg['dropout']))
            self.lstm=torch.nn.LSTM(n_nodes*h*len(spec['branch_inputs']),int(cfg['lstm_hidden']),batch_first=True)
            self.fc=torch.nn.Linear(int(cfg['lstm_hidden']),n_outputs)
        def encode(self,x,edges):

            expected_features=30 if self.scheme==6 else 2
            if x.ndim!=4 or x.shape[1]!=self.n_nodes or x.shape[-1]!=expected_features:
                raise ValueError(f'Input must have shape [B,N,T,{expected_features}], with fixed node order.')
            if len(edges)!=len(self.first): raise ValueError('The number of graphs does not match the scheme.')
            xt=x.permute(0,2,1,3)
            hs=[]
            for j,(first,second,edge) in enumerate(zip(self.first,self.second,edges)):
                inp=xt if self.scheme==6 else xt[...,spec['channels'][j]:spec['channels'][j]+1]
                h=self.dropout(torch.relu(first(inp,edge)))
                hs.append(self.dropout(torch.relu(second(h,edge))))
            states=torch.cat(hs,dim=-1)
            return states.reshape(states.shape[0],states.shape[1],-1)
        def forward(self,x,*edges):
            states,_=self.lstm(self.encode(x,edges))
            return self.fc(states[:,-1,:])
    return SchemeModel()

def fit_scalers(x,y,windows):
    y=np.asarray(y,float).reshape(len(x),-1)
    tr=windows.loc[windows.split=='train']
    used=np.zeros(len(x),bool)
    for a,b in zip(tr.input_start_index,tr.input_end_index): used[int(a):int(b)+1]=True
    if x.shape[-1]==30:
        lo=x[used].min(axis=0); hi=x[used].max(axis=0)
    elif x.shape[-1]==2:
        flat=x[used].reshape(-1,2); lo=flat.min(0); hi=flat.max(0)
    else: raise ValueError('Feature dimension must be 2 for grid inputs or 30 for station PCA inputs.')
    yt=y[tr.target_index.to_numpy(int)]; yl=yt.min(0); yh=yt.max(0)
    if np.any(yl==yh): raise ValueError('At least one training target is constant.')
    return {'input_min':lo.tolist(),'input_scale':np.where(hi>lo,hi-lo,1.).tolist(),
            'target_min':yl.tolist(),'target_scale':(yh-yl).tolist(),
            'input_fit_first_index':int(np.flatnonzero(used)[0]),'input_fit_last_index':int(np.flatnonzero(used)[-1]),
            'target_fit_n_samples':len(tr),'method':('training days only; per-station per-PC inputs; per-station targets; no clipping' if x.shape[-1]==30 else 'training input days only; per-channel inputs; per-output targets; no clipping')}

def scale_arrays(x,y,scalers):
    xx=(x-np.asarray(scalers['input_min']))/np.asarray(scalers['input_scale'])
    yy=(np.asarray(y).reshape(len(x),-1)-np.asarray(scalers['target_min']))/np.asarray(scalers['target_scale'])
    return xx.astype(np.float32),yy.astype(np.float32)

def make_scheme_loader(x,y,windows,label,cfg,shuffle=False):
    torch,_,_=torch_dependencies()
    rows=windows.loc[windows.split==label].reset_index(drop=True)
    class Dataset(torch.utils.data.Dataset):
        def __len__(self): return len(rows)
        def __getitem__(self,i):
            r=rows.iloc[i]
            return torch.from_numpy(numpy_window(x,r)),torch.from_numpy(np.asarray(y[int(r.target_index)],np.float32))
    generator=torch.Generator().manual_seed(int(cfg['seed']))
    return torch.utils.data.DataLoader(Dataset(),batch_size=int(cfg['batch_size']),shuffle=shuffle,
                                      num_workers=0,drop_last=False,generator=generator),rows

def train_scheme(model,loaders,edges,cfg,device):

    torch,_,_=torch_dependencies()
    optimizer=torch.optim.Adam(model.parameters(),lr=float(cfg['learning_rate']),weight_decay=float(cfg['weight_decay']))
    scheduler=torch.optim.lr_scheduler.ReduceLROnPlateau(optimizer,factor=.5,patience=int(cfg['lr_patience']))
    best=float('inf'); best_epoch=0; wait=0; best_state=None
    for epoch in range(1,int(cfg['epochs'])+1):
        losses={}
        for label in ('train','val'):
            training=label=='train'; model.train(training); total=count=0
            with torch.set_grad_enabled(training):
                for xx,yy in loaders[label]:
                    xx,yy=xx.to(device),yy.to(device)
                    if training: optimizer.zero_grad(set_to_none=True)
                    pred=model(xx,*edges)
                    if pred.shape!=yy.shape: raise ValueError(f'Prediction/target shape mismatch: {pred.shape}/{yy.shape}')
                    loss=torch.mean((pred-yy)**2)
                    if not torch.isfinite(loss): raise ValueError('Training loss is nonfinite.')
                    if training: loss.backward(); optimizer.step()
                    total+=loss.item()*yy.numel(); count+=yy.numel()
            losses[label]=total/count
        scheduler.step(losses['val'])
        if losses['val']<best:
            best=losses['val']; best_epoch=epoch; wait=0
            best_state={k:v.detach().cpu().clone() for k,v in model.state_dict().items()}
        else: wait+=1
        if epoch==1 or epoch%10==0:
            print(f'epoch={epoch} train={losses["train"]:.6f} val={losses["val"]:.6f}',flush=True)
        if wait>=cfg['early_stop_patience']: break
    if best_state is None: raise ValueError('Training did not produce a valid model.')
    model.load_state_dict(best_state); model.eval()
    return {'best_epoch':best_epoch,'best_val_mse':best,'epochs_run':epoch}

def predict_loader(model,loader,edges,device,scalers):
    torch,_,_=torch_dependencies(); model.eval(); values=[]
    with torch.no_grad():
        for xx,_ in loader: values.append(model(xx.to(device),*edges).cpu().numpy())
    pred=np.concatenate(values).astype(float)*np.asarray(scalers['target_scale'])+np.asarray(scalers['target_min'])
    if not np.isfinite(pred).all(): raise ValueError('Predictions contain NaN/Inf.')
    return pred

def prepare_basin(basin,graphs):
    errors={}; edge_sets={}; stats={}
    dates,x,y,nodes=load_basin_inputs(basin)
    for prefix in ('surface','baseflow'):
        try:
            edge,_,st=build_flow_graph(basin['flow_dir_'+prefix],nodes,**graphs[prefix])
            stats[prefix]=st
            if st['edges']==0 or st['has_cycle'] or st['status_counts'].get('flow_cycle',0):
                raise ValueError('No valid inter-node edges, or a directed flow cycle was detected.')
            edge_sets[prefix]=edge
        except Exception as exc: errors[prefix]=str(exc)
    if x.shape!=(len(dates),len(nodes),2) or not np.isfinite(x).all() or not np.isfinite(y).all() or (x<0).any() or (y<0).any():
        raise ValueError('Invalid input dimensions, missing values, or negative values.')
    for edge in edge_sets.values():
        if edge.ndim!=2 or edge.shape[0]!=2 or (edge<0).any() or (edge>=len(nodes)).any():
            raise ValueError('Invalid graph indices.')
    return {'basin':basin,'dates':dates,'x':x,'y':y[:,None],'nodes':nodes,'edges':edge_sets,
            'graph_stats':stats,'graph_errors':errors}

def load_station_inputs(basin):

    surface,base=read_daily(basin['runoff_excel']),read_daily(basin['baseflow_excel'])
    if not surface.index.equals(base.index): raise ValueError('Scheme VI runoff and baseflow dates differ.')
    columns=[list(surface.columns[2:]),list(base.columns[2:])]
    channels=[surface.loc[:,columns[0]].to_numpy(float),base.loc[:,columns[1]].to_numpy(float)]
    y=surface.iloc[:,1].to_numpy(float)[:,None]
    base_y=base.iloc[:,1].to_numpy(float)[:,None]
    if not np.array_equal(y,base_y):
        raise ValueError('Scheme VI observed discharge differs between runoff and baseflow tables.')
    for values in channels+[y]:
        if not np.isfinite(values).all() or (values<0).any(): raise ValueError('Scheme VI inputs or targets contain missing, nonfinite, or negative values.')
    return {'basin':basin,'dates':surface.index,'raw_channels':channels,'feature_columns':columns,'y':y}

def read_station_edges(path):
    frame=pd.read_csv(path,dtype=str,keep_default_na=False)
    if list(frame.columns)!=['upstream','downstream']:
        raise ValueError('Station edge CSV must contain exactly upstream,downstream columns.')
    frame=frame.apply(lambda column:column.str.strip())
    if frame.empty or (frame=='').any().any():
        raise ValueError('Station edge CSV is empty or contains blank station names.')
    return list(frame.itertuples(index=False,name=None))

def station_topology(items,explicit=None):
    names=[d['basin']['name'] for d in items]
    if names!=STATION_ORDER: raise ValueError('Scheme VI station order must be ribu,yilinchang,dajin,danba,xiaojin.')
    pairs=[]
    for a,b in (STATION_EDGES if explicit is None else explicit):
        if a not in names or b not in names or a==b: raise ValueError('Station connections contain unknown names or self-connections.')
        pairs.append((names.index(a),names.index(b)))
    if len(pairs)!=len(set(pairs)): raise ValueError('Duplicate station connections.')
    edges=np.asarray(pairs,dtype=np.int64).reshape(-1,2).T
    if graph_has_cycle(edges,len(names)): raise ValueError('Station connections contain a cycle.')
    if any(sum(a==i for a,b in pairs)>1 for i in range(len(names))): raise ValueError('Each station may have at most one immediate downstream station.')
    adjacent=[set() for _ in names]
    for a,b in pairs: adjacent[a].add(b); adjacent[b].add(a)
    seen={0}; stack=[0]
    while stack:
        for j in adjacent[stack.pop()]-seen: seen.add(j); stack.append(j)
    if len(seen)!=len(names) or len(pairs)!=len(names)-1: raise ValueError('The station graph must be a connected drainage tree.')
    return edges,{'rule':'user_supplied_station_connectivity','station_order':names,
                  'edges':[{'upstream':names[a],'downstream':names[b]} for a,b in pairs]}

def training_input_mask(n_days,windows):
    mask=np.zeros(n_days,dtype=bool)
    for a,b in windows.loc[windows.split=='train',['input_start_index','input_end_index']].itertuples(index=False,name=None):
        mask[int(a):int(b)+1]=True
    return mask

def fit_station_pca(values,mask,components=15):

    train=np.asarray(values,dtype=np.float64)[mask]
    if min(train.shape)<components or len(train)<components+1:
        raise ValueError(f'Reducing each station/channel to {components} components requires at least {components} original features and {components+1} training input days; current shape: {train.shape}. Dimensions are not reduced automatically or padded with zeros.')
    mean=train.mean(0); centered=train-mean
    _,singular,vt=np.linalg.svd(centered,full_matrices=False)
    basis=vt[:components].copy()

    signs=np.sign(basis[np.arange(components),np.argmax(np.abs(basis),axis=1)])
    basis*=np.where(signs==0,1.,signs)[:,None]
    variance=singular**2/(len(train)-1); total=float(variance.sum())
    if total<=0: raise ValueError('All training features are constant; PCA cannot be fitted.')
    state={'mean':mean.tolist(),'components':basis.tolist(),'explained_variance':variance[:components].tolist(),
           'explained_variance_ratio':(variance[:components]/total).tolist(),'n_components':components,
           'n_features_in':train.shape[1],'n_fit_days':len(train),'solver':'numpy_full_svd',
           'whiten':False,'pre_pca_standardization':False}
    return apply_station_pca(values,state),state

def apply_station_pca(values,state):
    values=np.asarray(values,dtype=np.float64)
    if values.ndim!=2 or values.shape[1]!=state['n_features_in']: raise ValueError('The PCA input feature count differs from training.')
    transformed=(values-np.asarray(state['mean']))@np.asarray(state['components']).T
    if not np.isfinite(transformed).all(): raise ValueError('PCA produced nonfinite values.')
    return transformed

def prepare_station_data(items,cfg,explicit=None):
    lookup={d['basin']['name']:d for d in items}
    if len(items)!=5 or set(lookup)!=set(STATION_ORDER): raise ValueError('Scheme VI requires inputs for all five configured stations.')
    items=[lookup[name] for name in STATION_ORDER]
    if not all(d['dates'].equals(items[0]['dates']) for d in items): raise ValueError('All five Scheme VI stations must have identical dates; row truncation is not allowed.')
    dates=items[0]['dates']; windows,split=make_windows(dates,cfg)
    mask=training_input_mask(len(dates),windows)
    edge,evidence=station_topology(items,explicit)
    values=[]; pca_states=[]
    for item in items:
        projections=[]; channels={}
        for index,channel in enumerate(('runoff','baseflow')):
            scores,state=fit_station_pca(item['raw_channels'][index],mask,STATION_PCA_DIM)
            state['feature_columns']=item['feature_columns'][index]
            projections.append(scores); channels[channel]=state
        values.append(np.concatenate(projections,axis=1))
        pca_states.append({'name':item['basin']['name'],'channels':channels})
    x=np.stack(values,axis=1)
    y=np.concatenate([d['y'] for d in items],axis=1)
    nodes=pd.DataFrame({'node_id':STATION_ORDER,'feature_layout':'runoff_PC01..15,baseflow_PC01..15'})
    preprocessing={'method':'separate_per_station_per_channel_PCA','stations':pca_states,
                   'fit_first_date':str(dates[np.flatnonzero(mask)[0]].date()),
                   'fit_last_date':str(dates[np.flatnonzero(mask)[-1]].date()),
                   'fit_n_days':int(mask.sum()),'split':split}
    return {'dates':dates,'x':x,'y':y,'nodes':nodes,'edges':{'station':edge},
            'graph_stats':{'station':evidence},'graph_errors':{},'pca_preprocessing':preprocessing,
            'output_names':STATION_ORDER.copy()}

def selected_edges(data,scheme):
    result=[]
    for graph in scheme_definition(scheme)['graphs']:
        if graph not in data['edges']: raise ValueError(f'{graph} graph unavailable: {data["graph_errors"].get(graph,"missing")}')
        result.append(data['edges'][graph])
    return result

def export_simple_results(results,out,name,scheme,make_plots=True):

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    out=Path(out); out.mkdir(parents=True,exist_ok=True)
    wb=Workbook(); summary=wb.active; summary.title='Metrics'
    summary.append(['Period','KGE','NSE','MAE (m3/s)','RMSE (m3/s)'])
    all_metrics=[]
    for label,title in (('train','Train'),('val','Validation'),('test','Test')):
        frame=results[label]
        m=calc_metrics(frame.obs.to_numpy(),frame.sim.to_numpy())
        summary.append([title]+[float(m[k]) for k in ('KGE','NSE','MAE','RMSE')])
        all_metrics.append({'split':label,**{k:m[k] for k in ('KGE','NSE','MAE','RMSE')}})
        ws=wb.create_sheet(title); ws.append(['Date','Observed (m3/s)','Simulated (m3/s)'])
        for date,obs,sim in frame[['date','obs','sim']].itertuples(index=False,name=None):
            ws.append([pd.Timestamp(date).to_pydatetime(),float(obs),float(sim)])
        for row in ws.iter_rows(min_row=2):
            row[0].number_format='yyyy-mm-dd'
            row[1].number_format=row[2].number_format='0.00'
    for ws in wb:
        ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]:
            cell.font=Font(name='Calibri',bold=True,color='FFFFFF')
            cell.fill=PatternFill('solid',fgColor='24485C')
        ws.row_dimensions[1].height=24
        for col in ('A','B','C','D','E'): ws.column_dimensions[col].width=21
    for row in summary.iter_rows(min_row=2,min_col=2):
        for cell in row: cell.number_format='0.00'
    wb.save(out/f'{name}_Scheme_{scheme}_results.xlsx')
    if not make_plots:
        return all_metrics
    frame=results['test']; m=all_metrics[-1]
    fig,ax=plt.subplots(figsize=(12,4))
    ax.plot(pd.to_datetime(frame.date),frame.obs,color='black',lw=.85,label='Observed')
    ax.plot(pd.to_datetime(frame.date),frame.sim,color='#258B91',lw=.85,label='Simulated')
    ax.set(xlabel='Date',ylabel='Streamflow (m³/s)',title=f'{name.capitalize()} — Scheme {scheme}')
    ax.text(.015,.97,f"KGE = {m['KGE']:.2f}   NSE = {m['NSE']:.2f}\nMAE = {m['MAE']:.2f} m³/s   RMSE = {m['RMSE']:.2f} m³/s",
            transform=ax.transAxes,va='top',fontsize=10,bbox=dict(facecolor='white',alpha=.9,edgecolor='none'))
    ax.margins(y=.24); ax.legend(loc='upper right',frameon=False)
    ax.spines[['top','right']].set_visible(False); fig.tight_layout()
    fig.savefig(out/f'{name}_Scheme_{scheme}_test.png',dpi=250); plt.close(fig)
    return all_metrics

def run_trial(data,scheme,cfg,out):
    torch,_,_=torch_dependencies(); validate_training_config(cfg)
    names=data.get('output_names',[data.get('basin',{}).get('name')])
    windows,split=make_windows(data['dates'],cfg)
    scalers=fit_scalers(data['x'],data['y'],windows)
    xx,yy=scale_arrays(data['x'],data['y'],scalers)
    set_seed(int(cfg['seed']),bool(cfg['deterministic']),torch)
    device=torch.device(('cuda' if torch.cuda.is_available() else 'cpu') if cfg['device']=='auto' else cfg['device'])
    if device.type=='cpu': torch.set_num_threads(1)
    model=build_scheme_model(data['x'].shape[1],cfg,scheme,len(names)).to(device)
    edges=tuple(torch.from_numpy(e).long().to(device) for e in selected_edges(data,scheme))
    loaders={s:make_scheme_loader(xx,yy,windows,s,cfg,s=='train')[0] for s in ('train','val')}
    training_summary=train_scheme(model,loaders,edges,cfg,device)
    out=Path(out); out.mkdir(parents=True,exist_ok=True)
    schema={'kind':'grid','nodes':data['nodes'].to_dict(orient='records')}
    if scheme==6:
        schema={'kind':'station_pca15','preprocessing':data['pca_preprocessing']}

    torch.save({
        'format':'VIC_TGCN_SCHEMES_V1', 'scheme':scheme,
        'model_state_dict':{k:v.detach().cpu().clone() for k,v in model.state_dict().items()},
        'model_config':dict(cfg), 'n_nodes':int(data['x'].shape[1]),
        'output_names':names, 'scalers':scalers, 'data_schema':schema,
        'edges':[e.detach().cpu() for e in edges], 'split':split,
        'training_summary':training_summary,
    },out/'reproduction_checkpoint.pt')
    results={name:{} for name in names}
    for label in ('train','val','test'):
        loader,rows=make_scheme_loader(xx,yy,windows,label,cfg)
        pred=predict_loader(model,loader,edges,device,scalers)
        obs=data['y'][rows.target_index.to_numpy(int)]
        for j,name in enumerate(names):
            results[name][label]=pd.DataFrame({'date':rows.target_date,'obs':obs[:,j],'sim':pred[:,j]})
    records=[]
    for name in names:
        for m in export_simple_results(results[name],out,name,scheme,cfg['make_plots']):
            records.append({'basin':name,'scheme':scheme,'seed':cfg['seed'],**m})
    return pd.DataFrame(records)

def predict_saved(checkpoint_path,basins,out,device_name='auto',target_start=None,target_end=None):

    torch,_,_=torch_dependencies()
    cp=torch.load(checkpoint_path,map_location='cpu',weights_only=True)
    if cp.get('format')!='VIC_TGCN_SCHEMES_V1': raise ValueError('Use reproduction_checkpoint.pt generated by this script.')
    schema=cp['data_schema']; lookup={b['name']:b for b in basins}
    if schema['kind']=='station': raise ValueError('Legacy mean-feature Scheme VI checkpoints are incompatible with PCA15; retrain Scheme VI.')
    is_station=schema['kind']=='station_pca15'
    specs=schema['preprocessing']['stations'] if is_station else [{'name':cp['output_names'][0],'nodes':schema['nodes']}]
    values=[]; dates=None
    for spec in specs:
        if spec['name'] not in lookup: raise ValueError(f'Prediction configuration is missing {spec["name"]}.')
        basin=lookup[spec['name']]
        if is_station:
            sc=spec['channels']['runoff']['feature_columns']; bc=spec['channels']['baseflow']['feature_columns']
        else:
            nodes=pd.DataFrame(spec['nodes']); sc=nodes.surface_column.tolist(); bc=nodes.baseflow_column.tolist()
        s,xs=read_prediction_table(basin['runoff_excel'],sc)
        b,xb=read_prediction_table(basin['baseflow_excel'],bc)
        if not s.index.equals(b.index) or (dates is not None and not dates.equals(s.index)): raise ValueError('Prediction input dates differ.')
        dates=s.index
        if is_station:
            values.append(np.concatenate([apply_station_pca(xs,spec['channels']['runoff']),
                                          apply_station_pca(xb,spec['channels']['baseflow'])],axis=1))
        else: values.append(np.stack([xs,xb],-1))
    x=np.stack(values,axis=1) if is_station else values[0]
    cfg=cp['model_config']; scaler=cp['scalers']
    x=((x-np.asarray(scaler['input_min']))/np.asarray(scaler['input_scale'])).astype(np.float32)
    if not np.isfinite(x).all(): raise ValueError('Invalid normalized prediction inputs.')
    window,lead=int(cfg['window']),int(cfg['target_lead_days'])
    ends=np.arange(window-1,len(x)); targets=dates[ends]+pd.Timedelta(days=lead)
    keep=np.ones(len(ends),bool)
    if target_start: keep&=targets>=pd.Timestamp(target_start)
    if target_end: keep&=targets<=pd.Timestamp(target_end)
    ends,targets=ends[keep],targets[keep]
    if not len(ends): raise ValueError('No complete prediction windows.')
    set_seed(int(cfg['seed']),bool(cfg['deterministic']),torch)
    device=torch.device(('cuda' if torch.cuda.is_available() else 'cpu') if device_name=='auto' else device_name)
    if device.type=='cpu': torch.set_num_threads(1)
    model=build_scheme_model(cp['n_nodes'],cfg,cp['scheme'],len(cp['output_names'])).to(device)
    model.load_state_dict(cp['model_state_dict'],strict=True); model.eval()
    edges=tuple(e.to(device) for e in cp['edges']); output=[]
    with torch.no_grad():
        for first in range(0,len(ends),cfg['batch_size']):
            batch=np.stack([x[e-window+1:e+1].transpose(1,0,2) for e in ends[first:first+cfg['batch_size']]])
            output.append(model(torch.from_numpy(np.ascontiguousarray(batch)).to(device),*edges).cpu().numpy())
    pred=np.concatenate(output).astype(float)*np.asarray(scaler['target_scale'])+np.asarray(scaler['target_min'])
    if not np.isfinite(pred).all(): raise ValueError('Independent prediction produced nonfinite values.')
    out=Path(out); out.mkdir(parents=True,exist_ok=True)
    for j,name in enumerate(cp['output_names']):
        pd.DataFrame({'date':targets.strftime('%Y-%m-%d'),'input_start_date':dates[ends-window+1].strftime('%Y-%m-%d'),
                      'input_end_date':dates[ends].strftime('%Y-%m-%d'),'sim':pred[:,j]}).to_csv(out/f'{name}_prediction.csv',index=False,encoding='utf-8-sig')
    write_json(out/'prediction_manifest.json',{'checkpoint':str(checkpoint_path),'sha256':file_sha256(checkpoint_path),
               'input_hashes':{spec['name']:{k:file_sha256(lookup[spec['name']][k]) for k in ('runoff_excel','baseflow_excel')} for spec in specs},
               'scheme':cp['scheme'],'n_samples':len(ends),'target_lead_days':lead,'scaler_refit':False})
    print(f'Independent prediction completed: {out}',flush=True)
    return pred,targets

def self_test():

    import tempfile
    torch,_,_ = torch_dependencies()
    torch.set_num_threads(1)
    set_seed(42,True,torch)
    nodes = pd.DataFrame({"node_id":["center","north","east"],
                          "lat":[1.5,2.5,1.5],"lon":[1.5,1.5,2.5]})
    with tempfile.TemporaryDirectory() as td:
        path = Path(td)/"known_directions.txt"
        for code,target in ((1,1),(3,2)):
            path.write_text("ncols 3\nnrows 3\nxllcorner 0\nyllcorner 0\ncellsize 1\nNODATA_value 255\n0 0 0\n0 "+str(code)+" 0\n0 0 0\n",encoding="utf-8")
            edge,_,_ = build_flow_graph(path,nodes)
            assert edge.tolist()==[[0],[target]], edge
    dates=pd.date_range("2009-01-01",periods=80)
    cfg=dict(DEFAULT_TRAINING,gcn_hidden=2,lstm_hidden=4,dropout=0.0,window=3,make_plots=False)
    windows,_=make_windows(dates,cfg)
    assert (windows.input_end_index < windows.target_index).all()
    obs=np.array([1.,2.,4.,8.]); pred=np.array([1.2,2.1,3.6,7.7])
    metrics=calc_metrics(obs,pred)
    assert np.isclose(metrics["RMSE"]**2/(1-metrics["NSE"]),np.var(obs))
    edges=torch.tensor([[0,1],[1,2]],dtype=torch.long)
    for scheme in (2,3,4,5,6):
        n=5 if scheme==6 else 3
        features=30 if scheme==6 else 2
        model=build_scheme_model(n,cfg,scheme,n_outputs=5 if scheme==6 else 1).eval()
        x=torch.arange(1,1+2*n*4*features,dtype=torch.float32).reshape(2,n,4,features)/100
        ee=torch.tensor([[0,1,2,4],[2,2,3,3]],dtype=torch.long) if scheme==6 else edges
        spec=scheme_definition(scheme); graph_args=tuple(ee for _ in spec['branch_inputs'])
        actual=model.encode(x,graph_args)
        expected=[]
        for sample in x:
            times=[]
            for t in range(4):
                hs=[]
                for j,(first,second,edge) in enumerate(zip(model.first,model.second,graph_args)):
                    inp=sample[:,t,:] if scheme==6 else sample[:,t,spec['channels'][j]:spec['channels'][j]+1]
                    h=model.dropout(torch.relu(first(inp,edge)))
                    hs.append(model.dropout(torch.relu(second(h,edge))))
                times.append(torch.cat(hs,-1).flatten())
            expected.append(torch.stack(times))
        torch.testing.assert_close(actual,torch.stack(expected),rtol=1e-5,atol=1e-6)
        output=model(x,*graph_args)
        assert output.shape==(2,5 if scheme==6 else 1) and torch.isfinite(output).all()
        output.square().mean().backward()
    print("PASS: Schemes II-VI physical time ordering and forward/backward; north/east D8; date causality; NSE/RMSE identity.")

def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--config',help='JSON configuration; default: packaged Danba data.')
    parser.add_argument('--preflight',action='store_true',help='Validate inputs, graphs and training-only scaling without training.')
    parser.add_argument('--self-test',action='store_true',help='Run small synthetic implementation tests; no scientific results.')
    parser.add_argument('--schemes',nargs='+',type=int,choices=range(1,7))
    parser.add_argument('--basin',nargs='+')
    parser.add_argument('--seeds',nargs='+',type=int)
    parser.add_argument('--epochs',type=int)
    parser.add_argument('--device')
    parser.add_argument('--no-plots',action='store_true')
    parser.add_argument('--output')
    parser.add_argument('--mode',choices=['train','predict'],default='train')
    parser.add_argument('--checkpoint')
    parser.add_argument('--target-start'); parser.add_argument('--target-end')
    args=parser.parse_args()
    if args.self_test:
        self_test(); return
    if not args.preflight:
        torch_dependencies()[0].set_num_threads(1)
    basins=copy.deepcopy(BASINS); cfg=copy.deepcopy(TRAINING); graphs=copy.deepcopy(GRAPH_CONFIG)
    station_edges=STATION_EDGES; schemes=args.schemes or SCHEMES; seeds=args.seeds or SEEDS
    if args.config:
        path=Path(args.config).resolve(); conf=json.loads(path.read_text(encoding='utf-8-sig'))
        unknown=set(conf.get('training',{}))-set(DEFAULT_TRAINING)
        if unknown: raise ValueError(f'Unknown training settings: {sorted(unknown)}')
        basins=conf['basins']; cfg.update(conf.get('training',{})); graphs=merged_graph_config(conf.get('graphs',{}))
        station_edges=conf.get('station_edges',station_edges)
        if conf.get('station_edges_file'):
            if 'station_edges' in conf:
                raise ValueError('Specify station_edges_file or station_edges, not both.')
            edge_path=Path(conf['station_edges_file'])
            if not edge_path.is_absolute(): edge_path=path.parent/edge_path
            station_edges=read_station_edges(edge_path)
        if not args.schemes: schemes=conf.get('schemes',schemes)
        if not args.seeds: seeds=conf.get('seeds',seeds)
        for b in basins:
            for field in ('runoff_excel','baseflow_excel','node_table','flow_dir_surface','flow_dir_baseflow','output_dir'):
                if b.get(field) and not Path(b[field]).is_absolute(): b[field]=str(path.parent/b[field])
    if args.epochs is not None: cfg['epochs']=args.epochs
    if args.device: cfg['device']=args.device
    if args.no_plots: cfg['make_plots']=False
    if len(set(b['name'] for b in basins))!=len(basins): raise ValueError('Duplicate basin names.')
    if len(set(schemes))!=len(schemes) or any(s not in range(1,7) for s in schemes): raise ValueError('Invalid or duplicate scheme IDs.')
    if len(set(seeds))!=len(seeds) or any(not isinstance(s,int) or s<0 for s in seeds): raise ValueError('Invalid or duplicate seeds.')
    chosen=args.basin or SELECTED_BASINS
    if chosen and not set(chosen).issubset({b['name'] for b in basins}): raise ValueError('Basin name not found in the configuration.')
    selected=[b for b in basins if not chosen or b['name'] in chosen]
    if args.mode=='predict':
        if not args.checkpoint: raise ValueError('Prediction mode requires --checkpoint.')
        out=args.output or str(Path(args.checkpoint).parent/'independent_prediction'/datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%S_%fZ'))
        predict_saved(args.checkpoint,basins,out,cfg['device'],args.target_start,args.target_end); return
    validate_training_config(cfg)
    if 6 in schemes:
        required = set(STATION_ORDER)
        if set(b['name'] for b in basins) != required:
            raise ValueError('Scheme VI requires all five station inputs. Danba alone is not Scheme VI. See configs/scheme_VI_five_stations.template.json.')
    if args.preflight:
        report = {}
        for basin in selected if any(s != 6 for s in schemes) else []:
            data = prepare_basin(basin, graphs)
            for scheme in schemes:
                if scheme != 6: selected_edges(data, scheme)
            windows, split = make_windows(data['dates'], cfg)
            fit_scalers(data['x'], data['y'], windows)
            report[basin['name']] = {'split':split,'samples':windows.split.value_counts().to_dict(),
                                    'nodes':len(data['nodes']),'graphs':{k:v.shape[1] for k,v in data['edges'].items()}}
        if 6 in schemes:
            data = prepare_station_data([load_station_inputs(b) for b in basins],cfg,station_edges)
            windows, split = make_windows(data['dates'],cfg)
            fit_scalers(data['x'],data['y'],windows)
            report['scheme_VI'] = {'shape':list(data['x'].shape),'split':split}
        print(json.dumps(report,ensure_ascii=False,indent=2))
        print('Preflight passed. No model was trained.')
        return
    parent=Path(args.output or COMPARISON_OUTPUT or (Path(basins[0]['output_dir']).parent/'Schemes_II_VI_results'))
    root=parent/'runs'/datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%S_%fZ')
    root.mkdir(parents=True,exist_ok=False)
    print(f'Comparison output directory: {root}',flush=True)
    lock=root/'RUNNING.lock'
    try: fd=os.open(lock,os.O_CREAT|os.O_EXCL|os.O_WRONLY)
    except FileExistsError as exc: raise RuntimeError('This directory is locked; verify the previous process has ended before removing RUNNING.lock.') from exc
    os.write(fd,str(os.getpid()).encode()); os.close(fd)
    failures=[]; prepared={}
    try:

        needed=selected if any(s!=6 for s in schemes) else []
        for b in needed:
            try: prepared[b['name']]=prepare_basin(b,graphs)
            except Exception as exc:
                failures.append({'task':'prepare','basin':b['name'],'error':str(exc)})
                print(f'{b["name"]} input preparation failed: {exc}',flush=True)
        for b in selected:
            if b['name'] not in prepared: continue
            d=prepared[b['name']]
            for seed in seeds:
                trial_cfg=dict(cfg,seed=seed)
                for scheme in [s for s in schemes if s!=6]:
                    out=root/f'scheme_{scheme}'/b['name']/f'seed_{seed}'
                    try: run_trial(d,scheme,trial_cfg,out)
                    except Exception as exc:
                        failures.append({'task':f'scheme_{scheme}','basin':b['name'],'seed':seed,'error':str(exc)})
                        out.mkdir(parents=True,exist_ok=True); (out/'ERROR.txt').write_text(traceback.format_exc(),encoding='utf-8')
                        print(f'Scheme {scheme} {b["name"]} failed: {exc}',flush=True)
        if 6 in schemes:
            try:
                items=[load_station_inputs(b) for b in basins]
                station=prepare_station_data(items,cfg,station_edges)
                print('Scheme VI station connections: '+str(station['graph_stats']['station']['edges']),flush=True)
                for seed in seeds:
                    out=root/'scheme_6'/'joint_stations'/f'seed_{seed}'
                    try: run_trial(station,6,dict(cfg,seed=seed),out)
                    except Exception as exc:
                        failures.append({'task':'scheme_6','seed':seed,'error':str(exc)})
                        out.mkdir(parents=True,exist_ok=True); (out/'ERROR.txt').write_text(traceback.format_exc(),encoding='utf-8')
                        print(f'Scheme VI seed={seed} failed: {exc}',flush=True)
            except Exception as exc:
                failures.append({'task':'scheme_6_prepare','error':str(exc)})
                (root/'Scheme6_ERROR.txt').write_text(traceback.format_exc(),encoding='utf-8')
                print(f'Scheme VI preparation failed: {exc}',flush=True)
        print(f'Results: {root}',flush=True)
    finally: lock.unlink(missing_ok=True)
    if failures: raise SystemExit('Some tasks did not finish; see the corresponding ERROR.txt files. Completed results have been saved.')

if __name__=='__main__':
    main()
