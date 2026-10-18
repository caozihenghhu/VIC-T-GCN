"""VIC-T-GCN Scheme I: dual-graph streamflow prediction."""
from __future__ import annotations
import argparse
import copy
import hashlib
import io
import json
import os
import random
import sys
from datetime import datetime, timezone
from pathlib import Path
import numpy as np
import pandas as pd

DEFAULT_TRAINING = {
    "window": 12, "target_lead_days": 1,
    "train_fraction": 0.65, "val_fraction": 0.10,
    "train_end": None, "val_end": None, "boundary_context": "past",
    "gcn_hidden": 32, "lstm_hidden": 128, "dropout": 0.3,
    "learning_rate": 0.001, "weight_decay": 1e-4, "batch_size": 8,
    "epochs": 200, "early_stop_patience": 15, "lr_patience": 5,
    "seed": 42, "deterministic": True, "device": "auto", "make_plots": True,
}

PACKAGE_ROOT = Path(__file__).resolve().parent
RUN_MODE = "train_predict"
SELECTED_BASINS = ["danba"]
BASINS = [{
    "name": "danba",
    "runoff_excel": str(PACKAGE_ROOT / "data/danba/runoff.xlsx"),
    "baseflow_excel": str(PACKAGE_ROOT / "data/danba/baseflow.xlsx"),
    "node_table": str(PACKAGE_ROOT / "data/danba/nodes.xlsx"),
    "flow_dir_surface": str(PACKAGE_ROOT / "data/danba/surface_direction.txt"),
    "flow_dir_baseflow": str(PACKAGE_ROOT / "data/danba/baseflow_direction.txt"),
    "output_dir": str(PACKAGE_ROOT / "results/scheme_I"),
}]

PREDICT_OPTIONS = {
    "checkpoint": None,
    "runoff_excel": None,
    "baseflow_excel": None,
    "output_dir": None,
    "target_start": None,
    "target_end": None,
    "observation_column": None,
    "device": "auto",
    "make_plot": True,
}

DIRECTION_MAPS = {
    "north_clockwise_1to8": {1:(-1,0),2:(-1,1),3:(0,1),4:(1,1),5:(1,0),6:(1,-1),7:(0,-1),8:(-1,-1)},
    "east_clockwise_1to8": {1:(0,1),2:(1,1),3:(1,0),4:(1,-1),5:(0,-1),6:(-1,-1),7:(-1,0),8:(-1,1)},
    "esri_d8": {1:(0,1),2:(1,1),4:(1,0),8:(1,-1),16:(0,-1),32:(-1,-1),64:(-1,0),128:(-1,1)},
}

def write_json(path, data):
    Path(path).write_text(json.dumps(data, ensure_ascii=False, indent=2, allow_nan=False), encoding="utf-8")

def file_sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for block in iter(lambda: f.read(1024*1024), b""):
            h.update(block)
    return h.hexdigest()

def read_table(path):
    path = Path(path)
    df = pd.read_excel(path) if path.suffix.lower() in (".xlsx", ".xls") else pd.read_csv(path, encoding="utf-8-sig")
    df.columns = [str(c).strip() for c in df.columns]
    if df.columns.duplicated().any():
        raise ValueError(f"Duplicate column names: {path}")
    return df

def parse_dates(values):

    v = pd.Series(values)
    if v.isna().any():
        raise ValueError("The date column contains missing values.")
    text = v.astype(str).str.strip()
    if text.str.fullmatch(r"\d{8}").all():
        dates = pd.to_datetime(text, format="%Y%m%d", errors="raise")
    elif pd.api.types.is_numeric_dtype(v):
        if not ((v >= 20000) & (v <= 80000) & (v == np.floor(v))).all():
            raise ValueError("The numeric first column is not a recognized Excel date; provide valid dates.")
        dates = pd.to_datetime(v, unit="D", origin="1899-12-30")
    else:
        dates = pd.to_datetime(v, errors="raise")
    dates = pd.DatetimeIndex(dates)
    if dates.tz is not None or not dates.equals(dates.normalize()):
        raise ValueError("Dates must be timezone-naive daily dates without time components.")
    if dates.has_duplicates:
        raise ValueError("Duplicate dates.")
    return dates

def read_daily(path):
    df = read_table(path)
    if df.shape[1] < 3 or len(df) == 0:
        raise ValueError(f"{path} requires dates, outlet observations, and at least one node input column.")
    df.index = parse_dates(df.iloc[:, 0])
    df = df.sort_index()
    if not df.index.equals(pd.date_range(df.index[0], df.index[-1], freq="D")):
        raise ValueError(f"{path} has nonconsecutive dates; missing days must not be treated as consecutive days.")
    return df

def validate_node_mapping(frame):

    required = ["node_index","node_id","lat","lon","surface_column","baseflow_column"]
    if not set(required).issubset(frame.columns):
        raise ValueError(f"The node table must contain columns: {required}")
    if len(frame) == 0 or frame[required].isna().any().any():
        raise ValueError("The node table is empty or required fields are missing.")
    nodes = frame[required].copy()
    for key in ("node_id","surface_column","baseflow_column"):
        nodes[key] = nodes[key].map(lambda v: str(int(v)) if isinstance(v,(float,np.floating)) and np.isfinite(v) and v.is_integer() else str(v).strip())
        if nodes[key].eq("").any() or nodes[key].duplicated().any():
            raise ValueError(f"Node table field {key} contains missing or duplicate values.")
    index = pd.to_numeric(nodes.node_index,errors="raise").to_numpy(float)
    if not np.isfinite(index).all() or not np.array_equal(np.sort(index),np.arange(len(nodes))):
        raise ValueError("node_index must contain unique consecutive integers from 0 to N-1.")
    nodes["node_index"] = index.astype(int)
    for key in ("lat","lon"):
        nodes[key] = pd.to_numeric(nodes[key],errors="raise").astype(float)
    if not np.isfinite(nodes[["lat","lon"]].to_numpy()).all():
        raise ValueError("Node coordinates contain missing or infinite values.")
    if not (nodes.lat.between(-90,90).all() and nodes.lon.between(-180,180).all()):
        raise ValueError("Node coordinates must be EPSG:4326 longitude/latitude, not projected coordinates.")
    if nodes[["lat","lon"]].duplicated().any():
        raise ValueError("Multiple nodes share the same coordinates.")
    nodes["coordinate_source"] = "node_table"
    nodes["mapping_rule"] = "explicit_node_index_and_named_input_columns"
    return nodes.sort_values("node_index").reset_index(drop=True)

def load_basin_inputs(basin):
    s, b = read_daily(basin["runoff_excel"]), read_daily(basin["baseflow_excel"])
    if not s.index.equals(b.index):
        raise ValueError("Runoff and baseflow dates differ; row-based concatenation is not allowed.")
    node_path = Path(basin["node_table"])
    if node_path.suffix.lower() in (".xlsx",".xls"):
        frame = pd.read_excel(node_path,sheet_name="Nodes",
                              dtype={"node_id":str,"surface_column":str,"baseflow_column":str})
    else:
        frame = pd.read_csv(node_path,encoding="utf-8-sig",
                            dtype={"node_id":str,"surface_column":str,"baseflow_column":str})
    nodes = validate_node_mapping(frame)
    for label, table, key in (("runoff",s,"surface_column"),("baseflow",b,"baseflow_column")):
        actual = list(table.columns[2:])
        expected = nodes[key].tolist()
        if len(actual) != len(expected) or set(actual) != set(expected):
            raise ValueError(f"{label} node columns do not match the node table; positional matching is not allowed.")
    xs = s.loc[:,nodes.surface_column].to_numpy(dtype=np.float64)
    xb = b.loc[:,nodes.baseflow_column].to_numpy(dtype=np.float64)
    y = s.iloc[:,1].to_numpy(dtype=np.float64)
    base_y = b.iloc[:,1].to_numpy(dtype=np.float64)
    if not np.array_equal(y,base_y):
        raise ValueError("Observed discharge differs between the runoff and baseflow tables.")
    for label, a in [("runoff",xs),("baseflow",xb),("observed discharge",y)]:
        if not np.isfinite(a).all() or np.any(a < 0):
            raise ValueError(f"{label} contains missing, infinite, or negative values; full-period interpolation is not performed.")
    if np.array_equal(xs,xb):
        raise ValueError("Runoff and baseflow are identical; check whether the same file was loaded twice.")
    print(f"Node table loaded: {len(nodes)} nodes; ordered by node_index and matched to both inputs by column name.",flush=True)
    return s.index, np.stack([xs,xb],axis=-1), y, nodes

def make_windows(dates, cfg):

    n, window, lead = len(dates), int(cfg["window"]), int(cfg["target_lead_days"])
    if window < 1 or lead < 0:
        raise ValueError("window must be positive and target_lead_days must be nonnegative.")
    if bool(cfg.get("train_end")) != bool(cfg.get("val_end")):
        raise ValueError("Specify both train_end and val_end, or set both to null.")
    if cfg.get("train_end"):
        t1, t2 = pd.Timestamp(cfg["train_end"]), pd.Timestamp(cfg["val_end"])
        if t1 not in dates or t2 not in dates or t1 >= t2:
            raise ValueError("Split dates must exist and train_end must precede val_end.")
        b1, b2 = dates.get_loc(t1)+1, dates.get_loc(t2)+1
    else:
        f1, f2 = float(cfg["train_fraction"]), float(cfg["val_fraction"])
        if not (0 < f1 < 1 and 0 < f2 < 1 and f1+f2 < 1):
            raise ValueError("Invalid training/validation fractions.")
        b1, b2 = int(n*f1), int(n*(f1+f2))
    if not 0 < b1 < b2 < n:
        raise ValueError("Empty or invalid data split.")
    target = np.arange(window-1+lead, n)
    last, first = target-lead, target-lead-window+1
    split = np.where(target < b1, "train", np.where(target < b2, "val", "test"))
    keep = np.ones(len(target), dtype=bool)
    if cfg["boundary_context"] == "within_split":
        keep = first >= np.where(target < b1, 0, np.where(target < b2, b1, b2))
    elif cfg["boundary_context"] != "past":
        raise ValueError("boundary_context must be past or within_split.")
    table = pd.DataFrame({
        "target_index": target[keep], "input_start_index": first[keep], "input_end_index": last[keep],
        "split": split[keep], "target_date": dates[target[keep]].strftime("%Y-%m-%d"),
        "input_start_date": dates[first[keep]].strftime("%Y-%m-%d"),
        "input_end_date": dates[last[keep]].strftime("%Y-%m-%d")})
    for label in ("train", "val", "test"):
        if not (table.split == label).any():
            raise ValueError(f"{label} has no valid windows.")
    return table, {"train_end": str(dates[b1-1].date()), "val_end": str(dates[b2-1].date()),
                   "test_start": str(dates[b2].date()), "split_basis": "raw_daily_dates"}

def fit_training_scalers(x, y, windows):
    tr = windows.loc[windows.split == "train"]
    days = np.zeros(len(x), dtype=bool)
    for a,b in zip(tr.input_start_index, tr.input_end_index):
        days[int(a):int(b)+1] = True
    flat = x[days].reshape(-1, 2)
    lo, hi = flat.min(axis=0), flat.max(axis=0)
    yt = y[tr.target_index.to_numpy(dtype=int)]
    ymin, ymax = float(yt.min()), float(yt.max())
    if ymin == ymax:
        raise ValueError("Training outlet discharge is constant; the regression experiment cannot proceed.")
    return {"input_min": lo.tolist(), "input_max": hi.tolist(),
            "input_scale": np.where(hi > lo, hi-lo, 1.0).tolist(),
            "target_min": ymin, "target_max": ymax, "target_scale": ymax-ymin,
            "input_fit_first_index": int(np.flatnonzero(days)[0]),
            "input_fit_last_index": int(np.flatnonzero(days)[-1]),
            "input_fit_n_days": int(days.sum()), "target_fit_n_samples": len(tr),
            "method": "minmax_per_channel_over_training_days_and_nodes; no clipping"}

def transform_arrays(x, y, scalers):
    xx = (x-np.array(scalers["input_min"]))/np.array(scalers["input_scale"])
    yy = (y-scalers["target_min"])/scalers["target_scale"]
    return xx.astype(np.float32), yy.astype(np.float32)

def numpy_window(x, row):

    return np.ascontiguousarray(x[int(row.input_start_index):int(row.input_end_index)+1].transpose(1,0,2))

def read_ascii_grid(path):
    try:
        text = Path(path).read_text(encoding="utf-8-sig")
    except UnicodeDecodeError:
        text = Path(path).read_text(encoding="gbk")
    lines, header = text.splitlines(), {}
    known = {"ncols","nrows","xllcorner","yllcorner","xllcenter","yllcenter","cellsize","nodata_value"}
    for start,line in enumerate(lines):
        parts = line.split()
        if not parts:
            continue
        key = parts[0].lower()
        if key not in known:
            break
        if len(parts) != 2 or key in header:
            raise ValueError(f"Invalid ASCII header: {line}")
        header[key] = float(parts[1])
    else:
        raise ValueError("The direction file contains no raster data.")
    for k in ("nrows","ncols","cellsize","nodata_value"):
        if k not in header or not np.isfinite(header[k]):
            raise ValueError(f"The ASCII header is missing a valid field: {k}")
    nr,nc,cs = int(header["nrows"]),int(header["ncols"]),header["cellsize"]
    if nr != header["nrows"] or nc != header["ncols"] or min(nr,nc,cs) <= 0:
        raise ValueError("Invalid raster dimensions or cellsize.")
    origin = []
    for ax in ("x","y"):
        corner,center = ax+"llcorner",ax+"llcenter"
        if (corner in header) == (center in header):
            raise ValueError(f"Specify exactly one of {corner} or {center}.")
        origin.append(header[corner] if corner in header else header[center]-cs/2)
    grid = np.loadtxt(io.StringIO("\n".join(lines[start:])), ndmin=2)
    if grid.shape != (nr,nc) or not np.isfinite(grid).all() or not np.isfinite(origin).all():
        raise ValueError("Invalid raster shape or values.")
    if not np.all(grid == np.rint(grid)):
        raise ValueError("Raster directions are not integers and cannot be used as D8 codes.")
    return grid, origin[0], origin[1], cs, header["nodata_value"]

def resolve_direction_mapping(encoding):
    if isinstance(encoding,str) and encoding in DIRECTION_MAPS:
        return DIRECTION_MAPS[encoding]
    if isinstance(encoding,dict):
        mapping = {int(k):tuple(int(v) for v in d) for k,d in encoding.items()}
        if len(mapping) != 8 or set(mapping.values()) != set(DIRECTION_MAPS["esri_d8"].values()) or 0 in mapping:
            raise ValueError("Custom codes must map uniquely to the eight neighbors; 0 is reserved for no direction.")
        return mapping
    raise ValueError("Specify verified surface_encoding/baseflow_encoding; values 1-8 alone do not identify the convention.")

def graph_has_cycle(edges, n):

    nxt = dict(edges.T.tolist())
    done = set()
    for start in range(n):
        path, node = set(), start
        while node in nxt and node not in done:
            if node in path:
                return True
            path.add(node)
            node = nxt[node]
        done.update(path)
    return False

def build_flow_graph(path, nodes, encoding="north_clockwise_1to8", crs="EPSG:4326", edge_rule="adjacent"):
    if crs != "EPSG:4326":
        raise ValueError("The ASCII raster and lat/lon must use EPSG:4326; reproject projected rasters first.")
    mapping = resolve_direction_mapping(encoding)
    if edge_rule not in ("adjacent","trace"):
        raise ValueError("edge_rule must be adjacent or trace.")
    grid,xll,yll,cs,nodata = read_ascii_grid(path)
    nr,nc = grid.shape
    bad = sorted(set(np.unique(grid))-set(mapping)-{0,nodata})
    if bad:
        raise ValueError(f"Raster contains undefined direction codes: {bad[:12]}; check the D8 convention and depression codes.")
    r = np.floor((yll+nr*cs-nodes.lat.to_numpy(float))/cs).astype(int)
    c = np.floor((nodes.lon.to_numpy(float)-xll)/cs).astype(int)
    if np.any((r<0)|(r>=nr)|(c<0)|(c>=nc)):
        raise ValueError("Nodes fall outside the raster; check CRS, coordinate order, and spatial coverage.")
    rc = list(zip(r.tolist(),c.tolist()))
    if len(set(rc)) != len(rc):
        raise ValueError("Multiple nodes map to the same raster cell; check resolution rather than overwriting node IDs.")

    missing_direction = grid[r,c] == nodata
    lookup, records = {v:i for i,v in enumerate(rc)}, []
    for i,(r0,c0) in enumerate(rc):
        rr,cc,seen,steps,dst = r0,c0,{(r0,c0)},0,None
        while True:
            code = int(grid[rr,cc])
            if code == nodata:
                status = "source_nodata" if steps == 0 else "path_reaches_nodata"
                break
            if code not in mapping:
                status = "no_direction"
                break
            dr,dc = mapping[code]
            rr,cc,steps = rr+dr,cc+dc,steps+1
            if not (0<=rr<nr and 0<=cc<nc):
                status = "leaves_raster"
                break
            if (rr,cc) in seen:
                status = "flow_cycle"
                break
            if (rr,cc) in lookup:
                dst,status = lookup[(rr,cc)],"edge"
                break
            if edge_rule == "adjacent":
                status = "downstream_cell_has_no_node"
                break
            seen.add((rr,cc))
        records.append({"source_index":i,"source_node":str(nodes.iloc[i].node_id),
                        "source_row":r0,"source_col":c0,"source_code":int(grid[r0,c0]),
                        "source_lat":float(nodes.iloc[i].lat),"source_lon":float(nodes.iloc[i].lon),
                        "source_direction_missing":bool(missing_direction[i]),
                        "target_index":dst,"target_node":None if dst is None else str(nodes.iloc[dst].node_id),
                        "end_row":rr,"end_col":cc,"raster_steps":steps,"status":status})
    audit = pd.DataFrame(records)
    pairs = [(v["source_index"],v["target_index"]) for v in records if v["status"]=="edge"]
    edges = np.array(pairs,dtype=np.int64).T if pairs else np.empty((2,0),dtype=np.int64)
    stats = {"nodes":len(nodes),"edges":int(edges.shape[1]),"isolated_nodes":len(nodes)-len(set(edges.ravel())),
             "has_cycle":graph_has_cycle(edges,len(nodes)),"edge_rule":edge_rule,"cellsize":cs,
             "nodata_nodes":int(missing_direction.sum()),
             "nodata_fraction":float(missing_direction.mean()),
             "nodata_policy":"retain_nodes_no_unknown_outgoing_edges_keep_supported_incoming_edges",
             "routing_direction_coverage_complete":not bool(missing_direction.any()),
             "status_counts":{str(k):int(v) for k,v in audit.status.value_counts().items()},
             "encoding":{str(k):list(v) for k,v in mapping.items()},"coordinate_system_assumption":crs}
    return edges,audit,stats

def torch_dependencies():
    try:
        import torch
        from torch_geometric.nn import GCNConv
        import torch_geometric
    except ImportError as exc:
        raise RuntimeError("Training requires compatible torch and torch_geometric installations; install torch, torch-geometric, numpy, pandas, scipy, openpyxl, and matplotlib.") from exc
    return torch,GCNConv,torch_geometric

def build_model(n_nodes, cfg):
    torch,GCNConv,_ = torch_dependencies()
    class DualGraphTGCN(torch.nn.Module):
        def __init__(self):
            super().__init__()
            h = int(cfg["gcn_hidden"])
            self.n_nodes = n_nodes
            self.gcn_s1,self.gcn_s2 = GCNConv(1,h),GCNConv(h,h)
            self.gcn_b1,self.gcn_b2 = GCNConv(1,h),GCNConv(h,h)
            self.dropout = torch.nn.Dropout(float(cfg["dropout"]))
            self.lstm = torch.nn.LSTM(2*h*n_nodes,int(cfg["lstm_hidden"]),batch_first=True)
            self.fc = torch.nn.Linear(int(cfg["lstm_hidden"]),1)

        def encode_lstm_input(self,x,edge_s,edge_b):
            if x.dim()==3:
                x = x.unsqueeze(0)
            if x.dim()!=4 or x.shape[1]!=self.n_nodes or x.shape[-1]!=2:
                raise ValueError("Input must have shape [B,N,T,2] or [N,T,2], with N matching the node order.")
            batch_states = []
            for sample in x:
                seq = []
                for t in range(sample.shape[1]):

                    hs = self.dropout(torch.relu(self.gcn_s1(sample[:,t,0:1],edge_s)))
                    hs = self.dropout(torch.relu(self.gcn_s2(hs,edge_s)))
                    hb = self.dropout(torch.relu(self.gcn_b1(sample[:,t,1:2],edge_b)))
                    hb = self.dropout(torch.relu(self.gcn_b2(hb,edge_b)))
                    seq.append(torch.cat([hs,hb],dim=-1))
                states = torch.stack(seq,dim=0)
                batch_states.append(states.reshape(states.shape[0],-1))
            return torch.stack(batch_states,dim=0)

        def forward(self,x,edge_s,edge_b):
            states,_ = self.lstm(self.encode_lstm_input(x,edge_s,edge_b))
            return self.fc(states[:,-1,:]).squeeze(-1)
    return DualGraphTGCN()

def set_seed(seed, deterministic, torch):
    if deterministic:
        os.environ.setdefault("CUBLAS_WORKSPACE_CONFIG",":4096:8")
    random.seed(seed)
    np.random.seed(seed)
    torch.manual_seed(seed)
    if torch.cuda.is_available():
        torch.cuda.manual_seed_all(seed)
    torch.backends.cudnn.benchmark = False
    torch.backends.cudnn.deterministic = bool(deterministic)

    torch.use_deterministic_algorithms(bool(deterministic))

def make_loader(x,y,windows,split,cfg,torch,shuffle=False):
    rows = windows.loc[windows.split==split].reset_index(drop=True)
    class WindowDataset(torch.utils.data.Dataset):
        def __len__(self):
            return len(rows)
        def __getitem__(self,index):
            row = rows.iloc[index]
            return (torch.from_numpy(numpy_window(x,row)),
                    torch.tensor(float(y[int(row.target_index)]),dtype=torch.float32))
    generator = torch.Generator().manual_seed(int(cfg["seed"]))
    loader = torch.utils.data.DataLoader(WindowDataset(),batch_size=int(cfg["batch_size"]),
                                        shuffle=shuffle,num_workers=0,drop_last=False,generator=generator)
    return loader,rows

def train_model(model,train_loader,val_loader,edges,cfg,device):

    torch,_,_ = torch_dependencies()
    optimizer = torch.optim.Adam(model.parameters(),lr=float(cfg["learning_rate"]),weight_decay=float(cfg["weight_decay"]))
    scheduler = torch.optim.lr_scheduler.ReduceLROnPlateau(optimizer,factor=0.5,patience=int(cfg["lr_patience"]))
    criterion = torch.nn.MSELoss()
    best,best_epoch,wait,best_state = float("inf"),0,0,None
    for epoch in range(1,int(cfg["epochs"])+1):
        losses = {}
        for label,loader in (("train",train_loader),("val",val_loader)):
            is_train = label=="train"
            model.train(is_train)
            total,count = 0.0,0
            with torch.set_grad_enabled(is_train):
                for xx,yy in loader:
                    xx,yy = xx.to(device),yy.to(device).reshape(-1)
                    if is_train:
                        optimizer.zero_grad(set_to_none=True)
                    pred = model(xx,*edges)
                    if pred.shape!=yy.shape:
                        raise RuntimeError(f"Prediction/target shape mismatch: {pred.shape} vs {yy.shape}")
                    loss = criterion(pred,yy)
                    if not torch.isfinite(loss):
                        raise FloatingPointError("Loss is NaN/Inf; training stopped.")
                    if is_train:
                        loss.backward()
                        optimizer.step()
                    total += loss.item()*yy.numel()
                    count += yy.numel()
            losses[label] = total/count
        scheduler.step(losses["val"])
        if losses["val"]<best:
            best,best_epoch,wait = losses["val"],epoch,0
            best_state = {k:v.detach().cpu().clone() for k,v in model.state_dict().items()}
        else:
            wait += 1
        if epoch==1 or epoch%10==0:
            print(f"epoch={epoch} train={losses['train']:.6f} val={losses['val']:.6f}",flush=True)
        if wait>=int(cfg["early_stop_patience"]):
            break
    if best_state is None:
        raise RuntimeError("No valid model was obtained.")
    model.load_state_dict(best_state)
    model.eval()
    return {"best_epoch":best_epoch,"best_val_mse":best,"epochs_run":epoch}

def predict(model,loader,edges,device,scalers):
    torch,_,_ = torch_dependencies()
    model.eval()
    p = []
    with torch.no_grad():
        for xx,_ in loader:
            p.append(model(xx.to(device),*edges).cpu().numpy().reshape(-1))
    pred = np.concatenate(p).astype(float)*scalers["target_scale"]+scalers["target_min"]
    if not np.isfinite(pred).all():
        raise FloatingPointError("Predictions contain NaN/Inf.")
    return pred

def calc_metrics(obs,pred):
    obs,pred = np.asarray(obs,float).reshape(-1),np.asarray(pred,float).reshape(-1)
    if obs.shape!=pred.shape or not len(obs) or not np.isfinite([obs,pred]).all():
        raise ValueError("Metrics require finite observations and predictions of equal length, paired by day.")
    mse,var,mean = float(np.mean((pred-obs)**2)),float(np.var(obs)),float(np.mean(obs))
    nse = 1-mse/var if var>0 else np.nan
    kge = np.nan
    if len(obs)>1 and var>0 and mean!=0 and np.std(pred)>0:
        r = float(np.corrcoef(obs,pred)[0,1])
        kge = 1-np.sqrt((r-1)**2+(np.std(pred)/np.std(obs)-1)**2+(np.mean(pred)/mean-1)**2)
    return {"NSE":nse,"KGE":float(kge),"MAE":float(np.mean(np.abs(pred-obs))),
            "RMSE":float(np.sqrt(mse))}

def validate_training_config(cfg):
    for k in ("window","batch_size","epochs","early_stop_patience","gcn_hidden","lstm_hidden"):
        if not isinstance(cfg[k],int) or cfg[k]<1:
            raise ValueError(f"{k} must be a positive integer.")
    for k in ("target_lead_days","lr_patience","seed"):
        if not isinstance(cfg[k],int) or cfg[k]<0:
            raise ValueError(f"{k} must be a nonnegative integer.")
    if not (0<=cfg["dropout"]<1 and cfg["learning_rate"]>0 and cfg["weight_decay"]>=0):
        raise ValueError("Invalid dropout, learning_rate, or weight_decay.")

def export_simple_results(results,out,name,scheme,make_plots=True):

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    out=Path(out); out.mkdir(parents=True,exist_ok=True)
    wb=Workbook(); summary=wb.active; summary.title='Metrics'
    summary.append(['Period','KGE','NSE','MAE (m3/s)','RMSE (m3/s)'])
    all_metrics=[]
    for label,title in (('train','Train'),('val','Validation'),('test','Test')):
        frame=results[label]
        m=calc_metrics(frame.obs.to_numpy(),frame.sim.to_numpy())
        summary.append([title]+[float(m[k]) for k in ('KGE','NSE','MAE','RMSE')])
        all_metrics.append({'split':label,**{k:m[k] for k in ('KGE','NSE','MAE','RMSE')}})
        ws=wb.create_sheet(title); ws.append(['Date','Observed (m3/s)','Simulated (m3/s)'])
        for date,obs,sim in frame[['date','obs','sim']].itertuples(index=False,name=None):
            ws.append([pd.Timestamp(date).to_pydatetime(),float(obs),float(sim)])
        for row in ws.iter_rows(min_row=2):
            row[0].number_format='yyyy-mm-dd'
            row[1].number_format=row[2].number_format='0.00'
    for ws in wb:
        ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]:
            cell.font=Font(name='Calibri',bold=True,color='FFFFFF')
            cell.fill=PatternFill('solid',fgColor='24485C')
        ws.row_dimensions[1].height=24
        for col in ('A','B','C','D','E'): ws.column_dimensions[col].width=21
    for row in summary.iter_rows(min_row=2,min_col=2):
        for cell in row: cell.number_format='0.00'
    wb.save(out/f'{name}_Scheme_{scheme}_results.xlsx')
    if not make_plots:
        return all_metrics
    frame=results['test']; m=all_metrics[-1]
    fig,ax=plt.subplots(figsize=(12,4))
    ax.plot(pd.to_datetime(frame.date),frame.obs,color='black',lw=.85,label='Observed')
    ax.plot(pd.to_datetime(frame.date),frame.sim,color='#258B91',lw=.85,label='Simulated')
    ax.set(xlabel='Date',ylabel='Streamflow (m³/s)',title=f'{name.capitalize()} — Scheme {scheme}')
    ax.text(.015,.97,f"KGE = {m['KGE']:.2f}   NSE = {m['NSE']:.2f}\nMAE = {m['MAE']:.2f} m³/s   RMSE = {m['RMSE']:.2f} m³/s",
            transform=ax.transAxes,va='top',fontsize=10,bbox=dict(facecolor='white',alpha=.9,edgecolor='none'))
    ax.margins(y=.24); ax.legend(loc='upper right',frameon=False)
    ax.spines[['top','right']].set_visible(False); fig.tight_layout()
    fig.savefig(out/f'{name}_Scheme_{scheme}_test.png',dpi=250); plt.close(fig)
    return all_metrics

def process_basin(basin,cfg,preflight=False):
    validate_training_config(cfg)
    dates,x,y,nodes=load_basin_inputs(basin)
    windows,split=make_windows(dates,cfg)
    scalers=fit_training_scalers(x,y,windows)
    edges=[]
    for prefix in ('surface','baseflow'):
        edge,audit,stats=build_flow_graph(basin[f'flow_dir_{prefix}'],nodes)
        if stats['edges']==0 or stats['has_cycle'] or stats['status_counts'].get('flow_cycle',0):
            raise ValueError(f'{prefix}: empty graph or flow cycle')
        edges.append(edge)
    print(f"{basin['name']}: {split}; samples={windows.split.value_counts().to_dict()}",flush=True)
    if preflight: return
    torch,_,_=torch_dependencies()
    set_seed(int(cfg['seed']),bool(cfg['deterministic']),torch)
    device=torch.device(('cuda' if torch.cuda.is_available() else 'cpu') if cfg['device']=='auto' else cfg['device'])
    xx,yy=transform_arrays(x,y,scalers)
    train_loader,_=make_loader(xx,yy,windows,'train',cfg,torch,shuffle=True)
    val_loader,_=make_loader(xx,yy,windows,'val',cfg,torch,shuffle=False)
    model=build_model(len(nodes),cfg).to(device)
    edge_tensors=tuple(torch.tensor(e,dtype=torch.long,device=device) for e in edges)
    out_dir=Path(basin['output_dir'])/'runs'/datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%S_%fZ')
    training_summary = train_model(model,train_loader,val_loader,edge_tensors,cfg,device)
    out_dir.mkdir(parents=True,exist_ok=True)

    torch.save({
        'format': 'VIC_TGCN_SCHEME_I_INFERENCE_V1', 'basin_name': basin['name'],
        'model_state_dict': {k:v.detach().cpu().clone() for k,v in model.state_dict().items()},
        'model_config': dict(cfg), 'n_nodes': len(nodes),
        'node_order': nodes.to_dict(orient='records'), 'scalers': scalers,
        'edge_surface': edge_tensors[0].detach().cpu(),
        'edge_baseflow': edge_tensors[1].detach().cpu(),
        'training_summary': training_summary, 'split': split,
    }, out_dir/'reproduction_checkpoint.pt')
    results={}
    for label in ('train','val','test'):
        loader,rows=make_loader(xx,yy,windows,label,cfg,torch,shuffle=False)
        pred=predict(model,loader,edge_tensors,device,scalers)
        results[label]=pd.DataFrame({'date':rows.target_date,'obs':y[rows.target_index.to_numpy(int)],'sim':pred})
    export_simple_results(results,out_dir,basin['name'],'I',cfg['make_plots'])
    print(f'Completed: {out_dir}',flush=True)
    return out_dir

def read_prediction_table(path, feature_columns):

    df = read_table(path)
    if len(df)==0 or len(df.columns)<2:
        raise ValueError(f"Prediction input is empty or lacks node columns: {path}")
    date_column = df.columns[0]
    if date_column in feature_columns:
        raise ValueError("The first prediction-input column must contain dates, not node features.")
    missing = set(feature_columns)-set(df.columns)
    if missing:
        raise ValueError(f"{path} is missing training node columns: {sorted(missing)[:10]}; positional matching is not allowed.")
    df.index = parse_dates(df.iloc[:,0])
    df = df.sort_index()
    if not df.index.equals(pd.date_range(df.index[0],df.index[-1],freq="D")):
        raise ValueError(f"Nonconsecutive prediction input dates: {path}")
    features = df.loc[:,feature_columns].to_numpy(dtype=np.float64)
    if not np.isfinite(features).all() or np.any(features<0):
        raise ValueError(f"Prediction features contain missing, infinite, or negative values: {path}")
    return df,features

def predict_from_checkpoint(options):

    for key in ("checkpoint","runoff_excel","baseflow_excel"):
        if not options.get(key) or not Path(options[key]).is_file():
            raise ValueError(f"Independent prediction requires an existing file: PREDICT_OPTIONS['{key}']")
    torch,_,_ = torch_dependencies()
    checkpoint_path = Path(options["checkpoint"])
    checkpoint = torch.load(checkpoint_path,map_location="cpu",weights_only=True)
    keys = {"model_state_dict","model_config","n_nodes","node_order","scalers","edge_surface","edge_baseflow"}
    if not isinstance(checkpoint,dict) or not keys.issubset(checkpoint):
        raise ValueError("Use reproduction_checkpoint.pt; best_weights.pt alone lacks graphs and scaling parameters.")
    cfg,scalers = checkpoint["model_config"],checkpoint["scalers"]
    nodes = pd.DataFrame(checkpoint["node_order"])
    if len(nodes)!=checkpoint["n_nodes"]:
        raise ValueError("Checkpoint node order and node count are inconsistent.")
    sc,bc = nodes.surface_column.tolist(),nodes.baseflow_column.tolist()
    s,xs = read_prediction_table(options["runoff_excel"],sc)
    b,xb = read_prediction_table(options["baseflow_excel"],bc)
    if not s.index.equals(b.index):
        raise ValueError("Runoff and baseflow prediction dates must match.")
    dates = s.index
    window,lead = int(cfg["window"]),int(cfg["target_lead_days"])
    if len(dates)<window:
        raise ValueError(f"Provide at least {window} consecutive input days; available: {len(dates)} days.")

    ends = np.arange(window-1,len(dates),dtype=int)
    target_dates = dates[ends]+pd.Timedelta(days=lead)
    keep = np.ones(len(ends),dtype=bool)
    if options.get("target_start"):
        keep &= target_dates>=pd.Timestamp(options["target_start"])
    if options.get("target_end"):
        keep &= target_dates<=pd.Timestamp(options["target_end"])
    ends,target_dates = ends[keep],target_dates[keep]
    if not len(ends):
        raise ValueError("No complete windows in the target date range; check input history and target_start/end.")
    x = np.stack([xs,xb],axis=-1)

    xx = ((x-np.array(scalers["input_min"]))/np.array(scalers["input_scale"])).astype(np.float32)
    if not np.isfinite(xx).all():
        raise ValueError("Nonfinite values after applying training scaling parameters.")
    device_name = options.get("device","auto")
    device = torch.device(("cuda" if torch.cuda.is_available() else "cpu") if device_name=="auto" else device_name)
    set_seed(int(cfg["seed"]),bool(cfg["deterministic"]),torch)
    model = build_model(int(checkpoint["n_nodes"]),cfg).to(device)
    model.load_state_dict(checkpoint["model_state_dict"])
    model.eval()
    edges = (checkpoint["edge_surface"].to(device),checkpoint["edge_baseflow"].to(device))
    output = []
    with torch.no_grad():
        for i in range(0,len(ends),int(cfg["batch_size"])):
            batch_ends = ends[i:i+int(cfg["batch_size"])]
            batch = np.stack([xx[e-window+1:e+1].transpose(1,0,2) for e in batch_ends])
            output.append(model(torch.from_numpy(np.ascontiguousarray(batch)).to(device),*edges).cpu().numpy())
    sim = np.concatenate(output).astype(float)*scalers["target_scale"]+scalers["target_min"]
    if not np.isfinite(sim).all():
        raise FloatingPointError("Predictions contain NaN/Inf.")
    frame = pd.DataFrame({"date":target_dates.strftime("%Y-%m-%d"),
                          "input_start_date":dates[ends-window+1].strftime("%Y-%m-%d"),
                          "input_end_date":dates[ends].strftime("%Y-%m-%d"),"sim":sim})
    obs_column = options.get("observation_column")
    if obs_column:
        if obs_column not in s.columns or obs_column in sc or obs_column==s.columns[0]:
            raise ValueError("observation_column must identify a separate observed-discharge column in the runoff prediction table.")
        obs = pd.to_numeric(s[obs_column],errors="raise").astype(float)
        if np.isinf(obs).any() or (obs.dropna()<0).any():
            raise ValueError("Observed discharge contains infinite or negative values; leave missing observations blank.")

        frame["obs"] = obs.reindex(target_dates).to_numpy()
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S_%fZ")
    root = Path(options["output_dir"]) if options.get("output_dir") else checkpoint_path.parent/"predictions"
    out_dir = root/stamp
    out_dir.mkdir(parents=True,exist_ok=False)
    frame.to_csv(out_dir/"predictions.csv",index=False,encoding="utf-8-sig")
    frame.to_excel(out_dir/"predictions.xlsx",index=False)
    write_json(out_dir/"scalers_used.json",scalers)
    nodes.to_csv(out_dir/"node_order_used.csv",index=False,encoding="utf-8-sig")
    evaluated = 0
    if "obs" in frame:
        valid = frame.obs.notna()
        evaluated = int(valid.sum())
        if evaluated:
            metrics = calc_metrics(frame.loc[valid,"obs"],frame.loc[valid,"sim"])
            pd.DataFrame([metrics]).to_csv(out_dir/"prediction_metrics.csv",index=False)
    write_json(out_dir/"prediction_manifest.json",{
        "basin":checkpoint.get("basin_name","not_recorded_in_older_checkpoint"),
        "checkpoint":str(checkpoint_path),"checkpoint_sha256":file_sha256(checkpoint_path),
        "runoff_input":str(options["runoff_excel"]),"runoff_sha256":file_sha256(options["runoff_excel"]),
        "baseflow_input":str(options["baseflow_excel"]),"baseflow_sha256":file_sha256(options["baseflow_excel"]),
        "input_start":str(dates[0].date()),"input_end":str(dates[-1].date()),
        "target_start":frame.date.iloc[0],"target_end":frame.date.iloc[-1],"n_predictions":len(frame),
        "window":window,"target_lead_days":lead,"fitted_on_prediction_data":False,
        "observations_used_for_model_input":False,"n_evaluated_with_observations":evaluated,
        "device":str(device),"negative_predictions":int((sim<0).sum()),
        "implementation_sha256":file_sha256(__file__)})
    if options.get("make_plot",True):
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        fig,ax = plt.subplots(figsize=(12,4))
        ax.plot(target_dates,frame.sim,label="Predicted",lw=1)
        if "obs" in frame:
            ax.plot(target_dates,frame.obs,label="Observed",lw=1)
        ax.set(xlabel="Target date",ylabel="Discharge (training target units)")
        ax.legend()
        fig.tight_layout()
        fig.savefig(out_dir/"predictions.png",dpi=180)
        plt.close(fig)
    print(f"Independent prediction completed: {len(frame)} days; output: {out_dir}",flush=True)
    return out_dir

def run_demo(output_root=None):

    root = Path(output_root) if output_root else Path(__file__).resolve().parent/"DEMO_OUTPUT"
    root = root/datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S_%fZ")
    data_dir = root/"input"
    data_dir.mkdir(parents=True,exist_ok=False)
    dates = pd.date_range("2009-01-01",periods=100,freq="D")
    day = np.arange(len(dates),dtype=float)
    columns = ["1", "2", "3"]
    pd.DataFrame({"node_index":[0,1,2],"node_id":columns,"lat":[0.5,0.5,0.5],
                  "lon":[0.5,1.5,2.5],"surface_column":columns,"baseflow_column":columns}).to_csv(
                  data_dir/"nodes.csv",index=False)
    s = pd.DataFrame({"date":dates,"Q":20+day/3+np.sin(day)})
    b = s.copy()
    for i,name in enumerate(columns):
        s[name] = 2+i+np.sin(day/4)+day/100
        b[name] = 1+i/3+np.cos(day/15)/5
    s.to_excel(data_dir/"runoff_train.xlsx",index=False)
    b.to_excel(data_dir/"baseflow_train.xlsx",index=False)
    grid = data_dir/"flow_direction.asc"
    grid.write_text("ncols 3\nnrows 1\nxllcorner 0\nyllcorner 0\ncellsize 1\nNODATA_value -9999\n3 3 0\n",encoding="utf-8")
    basin = {"name":"synthetic_demo","runoff_excel":str(data_dir/"runoff_train.xlsx"),
             "baseflow_excel":str(data_dir/"baseflow_train.xlsx"),
             "flow_dir_surface":str(grid),"flow_dir_baseflow":str(grid),"output_dir":str(root/"training"),
             "node_table":str(data_dir/"nodes.csv")}
    cfg = copy.deepcopy(DEFAULT_TRAINING)
    cfg.update({"epochs":3,"gcn_hidden":4,"lstm_hidden":8,"device":"cpu"})
    torch,_,_ = torch_dependencies()
    torch.set_num_threads(1)
    print("Running a synthetic demonstration; its scores do not represent real-basin performance.",flush=True)
    trained = process_basin(basin,cfg)

    s.drop(columns="Q").iloc[-25:][["date"]+columns[::-1]].to_excel(data_dir/"runoff_new_without_Q.xlsx",index=False)
    b.drop(columns="Q").iloc[-25:].to_excel(data_dir/"baseflow_new_without_Q.xlsx",index=False)
    predicted = predict_from_checkpoint({"checkpoint":str(trained/"reproduction_checkpoint.pt"),
        "runoff_excel":str(data_dir/"runoff_new_without_Q.xlsx"),"baseflow_excel":str(data_dir/"baseflow_new_without_Q.xlsx"),
        "output_dir":str(root/"new_predictions"),"device":"cpu","make_plot":True})
    print(f"Demo inputs: {data_dir}\nTraining outputs: {trained}\nIndependent prediction outputs: {predicted}",flush=True)
    return root

def self_test():

    import tempfile
    torch,_,_ = torch_dependencies()
    torch.set_num_threads(1)
    set_seed(42,True,torch)
    nodes = pd.DataFrame({"node_id":["center","north","east"],
                          "lat":[1.5,2.5,1.5],"lon":[1.5,1.5,2.5]})
    with tempfile.TemporaryDirectory() as td:
        path = Path(td)/"known_directions.txt"
        for code,target in ((1,1),(3,2)):
            path.write_text("ncols 3\nnrows 3\nxllcorner 0\nyllcorner 0\ncellsize 1\nNODATA_value 255\n0 0 0\n0 "+str(code)+" 0\n0 0 0\n",encoding="utf-8")
            edge,_,_ = build_flow_graph(path,nodes)
            assert edge.tolist()==[[0],[target]], edge
    dates=pd.date_range("2009-01-01",periods=80)
    cfg=dict(DEFAULT_TRAINING,gcn_hidden=2,lstm_hidden=4,dropout=0.0,window=3,make_plots=False)
    windows,_=make_windows(dates,cfg)
    assert (windows.input_end_index < windows.target_index).all()
    obs=np.array([1.,2.,4.,8.]); pred=np.array([1.2,2.1,3.6,7.7])
    metrics=calc_metrics(obs,pred)
    assert np.isclose(metrics["RMSE"]**2/(1-metrics["NSE"]),np.var(obs))
    edges=torch.tensor([[0,1],[1,2]],dtype=torch.long)
    model=build_model(3,cfg).eval()
    x=torch.arange(1,1+2*3*4*2,dtype=torch.float32).reshape(2,3,4,2)/10
    actual=model.encode_lstm_input(x,edges,edges)
    expected=[]
    for sample in x:
        times=[]
        for t in range(4):
            hs=model.dropout(torch.relu(model.gcn_s1(sample[:,t,0:1],edges)))
            hs=model.dropout(torch.relu(model.gcn_s2(hs,edges)))
            hb=model.dropout(torch.relu(model.gcn_b1(sample[:,t,1:2],edges)))
            hb=model.dropout(torch.relu(model.gcn_b2(hb,edges)))
            times.append(torch.cat([hs,hb],-1).flatten())
        expected.append(torch.stack(times))
    torch.testing.assert_close(actual,torch.stack(expected),rtol=0,atol=0)
    output=model(x,edges,edges)
    assert output.shape==(2,) and torch.isfinite(output).all()
    output.square().mean().backward()
    print("PASS: north/east D8 edges; physical time ordering; forward/backward; date causality; NSE/RMSE identity.")

def main():
    print("VIC-T-GCN Scheme I; the bundled default configuration uses Danba example inputs.", flush=True)
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--mode",choices=["train_predict","predict","demo"],default=None)
    parser.add_argument("--config",default=None,help="Optional JSON overrides; defaults to the configuration in this file")
    parser.add_argument("--basin",nargs="+",help="Basin names, e.g. danba; defaults to SELECTED_BASINS")
    parser.add_argument("--self-test",action="store_true",help="Run synthetic implementation checks, not a scientific experiment")
    parser.add_argument("--preflight",action="store_true",help="Validate training inputs without training")
    parser.add_argument("--checkpoint",help="Path to reproduction_checkpoint.pt for independent prediction")
    parser.add_argument("--runoff",help="Runoff Excel/CSV input for independent prediction")
    parser.add_argument("--baseflow",help="Baseflow Excel/CSV input for independent prediction")
    parser.add_argument("--output-dir",help="Output directory for independent prediction or the demo")
    parser.add_argument("--target-start")
    parser.add_argument("--target-end")
    parser.add_argument("--obs-column",help="Optional observation column used only for evaluation after prediction")
    parser.add_argument("--device",help="auto/cpu/cuda")
    args = parser.parse_args()
    if args.self_test:
        self_test(); return
    if not args.preflight:
        torch_dependencies()[0].set_num_threads(1)
    mode = args.mode or RUN_MODE
    if mode=="demo":
        run_demo(args.output_dir)
        return
    if mode=="predict":
        options = copy.deepcopy(PREDICT_OPTIONS)
        overrides = {"checkpoint":args.checkpoint,"runoff_excel":args.runoff,"baseflow_excel":args.baseflow,
                     "output_dir":args.output_dir,"target_start":args.target_start,"target_end":args.target_end,
                     "observation_column":args.obs_column,"device":args.device}
        options.update({k:v for k,v in overrides.items() if v is not None})
        for key in ("checkpoint","runoff_excel","baseflow_excel","output_dir"):
            if options.get(key) and not Path(options[key]).is_absolute():
                options[key] = str(Path(__file__).resolve().parent/options[key])
        predict_from_checkpoint(options)
        return
    if mode!="train_predict":
        raise ValueError("RUN_MODE must be train_predict, predict, or demo.")
    if args.config:
        cp = Path(args.config).resolve()
        config = json.loads(cp.read_text(encoding="utf-8-sig"))
        cfg = copy.deepcopy(DEFAULT_TRAINING)
        unknown = set(config.get("training",{}))-set(cfg)
        if unknown:
            raise ValueError(f"Unknown training configuration fields: {unknown}")
        cfg.update(config.get("training",{}))
        basins,base_dir = config["basins"],cp.parent
    else:
        cfg,basins,base_dir = copy.deepcopy(DEFAULT_TRAINING),copy.deepcopy(BASINS),Path(__file__).resolve().parent
    if args.device:
        cfg["device"] = args.device
    validate_training_config(cfg)
    chosen = args.basin if args.basin is not None else SELECTED_BASINS
    names = [b["name"] for b in basins]
    if len(set(names))!=len(names):
        raise ValueError("Duplicate basin names.")
    if chosen and not set(chosen).issubset(names):
        raise ValueError(f"Unknown basin; configured basins: {names}")
    failures = []
    for basin in basins:
        if chosen and basin["name"] not in chosen:
            continue
        basin = copy.deepcopy(basin)
        for key in ("runoff_excel","baseflow_excel","flow_dir_surface","flow_dir_baseflow","node_table","output_dir"):
            if basin.get(key) and not Path(basin[key]).is_absolute():
                basin[key] = str(base_dir/basin[key])
        try:
            process_basin(basin,cfg,preflight=args.preflight)
        except Exception as exc:
            failures.append(basin["name"])
            print(f"{basin['name']} failed: {exc}",file=sys.stderr,flush=True)
    if failures:
        raise SystemExit(f"Incomplete basins: {', '.join(failures)}. See the error messages above.")

if __name__=="__main__":
    main()
